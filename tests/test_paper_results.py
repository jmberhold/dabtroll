"""Paper tables must preserve cohorts, missingness, sequential gates, and clocks."""

import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
import pandas as pd
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.paper_results import (
    build_paper_results, export_tables, load_paper_results, plot_metric_by_scenario,
    plot_timing, progress_alignment, publication_tables, representative_cases,
    plot_simulator_endpoint, plot_prompt_sensitivity, plot_node_family_performance,
    plot_sequential_completion, plot_episode_triptych, plot_case,
)
from scripts import build_btaudit2_visualization_dataset as dataset


def fixture():
    episodes, nodes, observations = [], [], []
    for name, human_times, model_times, cohort in (
        ("a", [0, 2, 4, 6], [0, None, 5, 7], "evaluation"),
        ("b", [None, None], [1, 2], "evaluation"),
        ("c", [1, 2], None, "evaluation"),
        ("tune", [1, 2], [1, 2], "tuning"),
        ("unreviewed", [None, None], [1, 2], "evaluation"),
    ):
        row = {
            "episode_dir": name, "scenario": "Task", "cohort": cohort,
            "human_reviewed": name != "unreviewed", "human_comparable": name != "unreviewed",
            "review_basis": "fixture", "workbook_available": True,
            "sim_success_ever": name != "b", "sim_success_final": False,
            "required_nodes": len(human_times), "execution_nodes": len(human_times),
            "structurally_incomplete_nodes": 0, "bt_semantic_rating": None,
            "episode_path": "/fixture/" + name,
            "reviewer_progress_points_json": "[]",
            "progress_review_total_rows": 0, "progress_review_scored_rows": 0,
            "progress_review_duration_s": float("nan"),
        }
        episodes.append(row)
        for slot, human_time in enumerate(human_times, 1):
            nodes.append({
                **row, "node_id": f"node_action_{slot}", "node_position": slot,
                "human_node_available": True, "human_complete": human_time is not None,
                "human_completion_s": human_time, "reviewer_notes": "",
            })
        if model_times is not None:
            for slot, model_time in enumerate(model_times, 1):
                observations.append({
                    "episode_dir": name, "model_name": "qwen38",
                    "canonical_node_id": f"node_action_{slot}",
                    "model_status_norm": "complete" if model_time is not None else "running",
                    "video_time_s": model_time, "model_latency_s": 1,
                    "model_progress": slot / len(model_times),
                    "source": "mission_engine_response", "review_row": slot,
                    "machine_frame": f"frame_{slot:05}", "model_notes": "evidence",
                })
    return pd.DataFrame(episodes), pd.DataFrame(nodes), pd.DataFrame(observations)


class PaperResultsTests(unittest.TestCase):
    def test_incomplete_reference_criteria_do_not_shorten_full_mission_contract(self):
        episodes, nodes, obs = fixture()
        nodes = nodes[~(nodes.episode_dir.eq("b") & nodes.node_position.eq(1))]
        episodes.loc[episodes.episode_dir.eq("b"), "required_nodes"] = 1
        obs["node_position"] = obs["canonical_node_id"].str.rsplit("_", n=1).str[-1].astype(int)
        obs["model_node_count"] = obs["episode_dir"].map(
            episodes.set_index("episode_dir")["execution_nodes"])
        result = build_paper_results(episodes, nodes, obs)
        remaining = result.node_model.query("episode_dir == 'b'").iloc[0]
        self.assertEqual(remaining.node_position, 2)
        self.assertTrue(remaining.position_comparable)
        mission = result.episode_model.query("episode_dir == 'b'").iloc[0]
        self.assertTrue(mission.model_complete)
        self.assertFalse(mission.human_comparable)
        comparison = result.tables["node_position_reconciliation"].query("episode_dir == 'b'").iloc[0]
        self.assertEqual(comparison.reference_node_count, 2)
        self.assertEqual(comparison.model_node_count, 2)
        self.assertFalse(comparison.node_count_mismatch)

    def test_condition_slots_remain_distinct_and_unused_slots_empty(self):
        spec = dataset.LUNA6_SPEC
        observations = dataset._standardize_observations(pd.DataFrame({
            "node_id": ["node_condition_first", "node_condition_second"],
            "model_status": ["complete", "complete"],
            "video_time_m:ss": ["0:01", "0:02"],
        }), "luna6_v5", spec, {})
        observations["node_position"] = [1, 2]
        observations["model_node_count"] = 2
        record = {
            "required_nodes_json": json.dumps([
                {"node_id": "node_condition_first", "node_position": 1},
                {"node_id": "node_condition_second", "node_position": 2},
            ]),
            "node_1_node_id": "node_condition_first",
            "node_2_node_id": "node_condition_second",
        }
        dataset.nqcu._add_molmo_episode_fields(record, observations, "luna6", "luna6")
        self.assertEqual(record["node_1_luna6_model_completion_time"], "0:01")
        self.assertEqual(record["node_2_luna6_model_completion_time"], "0:02")
        for slot in (3, 4):
            self.assertEqual(record[f"node_{slot}_luna6_observation_count"], 0)
            self.assertFalse(record[f"node_{slot}_luna6_position_comparable"])

    def test_missing_models_keep_numeric_contract_counts(self):
        result = build_paper_results(*fixture())
        self.assertTrue(pd.api.types.is_numeric_dtype(result.node_model["model_node_count"]))
        episodes, nodes, observations = fixture()
        empty = build_paper_results(episodes, nodes, observations.iloc[:0])
        self.assertTrue(pd.api.types.is_numeric_dtype(empty.node_model["model_node_count"]))
        self.assertTrue(empty.tables["node_position_reconciliation"].empty)

    def test_ordinal_renaming_preserves_metrics_and_case_evidence(self):
        episodes, nodes, obs = fixture()
        obs["node_position"] = obs["canonical_node_id"].str.rsplit("_", n=1).str[-1].astype(int)
        obs["canonical_node_id"] = "different_" + obs["node_position"].astype(str)
        result = build_paper_results(episodes, nodes, obs)
        row = result.tables["node_metrics"].set_index("model_key").loc["qwen38"]
        self.assertEqual((row.n, row.tp, row.fp, row.fn), (6, 3, 2, 1))
        self.assertTrue(representative_cases(result)["model_notes"].eq("evidence").all())
        self.assertTrue(result.observations.canonical_node_id.str.startswith("different_").all())

    def test_extra_reference_positions_are_not_false_incomplete_or_whole_mission(self):
        episodes, nodes, obs = fixture()
        obs["node_position"] = obs["canonical_node_id"].str.rsplit("_", n=1).str[-1].astype(int)
        obs["model_node_count"] = obs["episode_dir"].map(episodes.set_index("episode_dir")["required_nodes"])
        obs.loc[obs.episode_dir.eq("a"), "model_node_count"] = 2
        obs = obs[~(obs.episode_dir.eq("a") & obs.node_position.gt(2))]
        result = build_paper_results(episodes, nodes, obs)
        row = result.tables["node_metrics"].set_index("model_key").loc["qwen38"]
        self.assertEqual((row.n, row.tp, row.fp, row.fn), (4, 1, 2, 1))
        self.assertEqual(row.not_comparable_reference_positions, 2)
        mission = result.tables["mission_metrics"].iloc[0]
        self.assertEqual((mission.n, mission.fp, mission.fn), (1, 1, 0))
        extra = result.node_model.query("episode_dir == 'a' and node_position > 2")
        self.assertFalse(extra.human_comparable.any())

    def test_extra_model_suffix_is_not_a_completed_whole_mission(self):
        episodes, nodes, obs = fixture()
        obs["node_position"] = obs["canonical_node_id"].str.rsplit("_", n=1).str[-1].astype(int)
        obs["model_node_count"] = obs["episode_dir"].map(episodes.set_index("episode_dir")["required_nodes"])
        obs.loc[obs.episode_dir.eq("b"), "model_node_count"] = 3
        result = build_paper_results(episodes, nodes, obs)
        row = result.episode_model.query("episode_dir == 'b'").iloc[0]
        self.assertFalse(row.model_complete)
        self.assertFalse(row.human_comparable)
        comparison = result.tables["node_position_reconciliation"].query("episode_dir == 'b'").iloc[0]
        self.assertEqual(comparison.model_only_positions, 1)
        self.assertEqual(comparison.shared_positions, 2)

    def test_prompt_error_comparison_requires_common_count_compatible_cohort(self):
        episodes, nodes, obs = fixture()
        obs["node_position"] = obs["canonical_node_id"].str.rsplit("_", n=1).str[-1].astype(int)
        obs["model_node_count"] = obs["episode_dir"].map(episodes.set_index("episode_dir")["required_nodes"])
        left = obs.assign(model_name="qwen35_qwen3_prompt")
        right = obs.assign(model_name="qwen35")
        right.loc[right.episode_dir.eq("a"), "model_node_count"] = 3
        right = right[~(right.episode_dir.eq("a") & right.node_position.eq(4))]
        result = build_paper_results(episodes, nodes, pd.concat([left, right]))
        errors = result.tables["prompt_error_tradeoff"]
        self.assertEqual(errors.n.tolist(), [1, 1])

    def test_offline_analysis_imports_without_replay_runtime(self):
        subprocess.run(
            [sys.executable, "-c",
             "import sys; from scripts import paper_results; "
             "assert 'scripts.replay_qwen35_bt_eval' not in sys.modules; "
             "assert 'dabtroll_bt_pipeline' not in sys.modules"],
            cwd=Path(__file__).resolve().parents[1], check=True,
        )

    def test_cohorts_missing_model_and_zero_time(self):
        result = build_paper_results(*fixture())
        row = result.tables["mission_metrics"].iloc[0]
        self.assertEqual((row.n, row.tp, row.fp, row.tn, row.fn), (2, 0, 1, 0, 1))
        nodes = result.tables["node_metrics"].set_index("model_key").loc["qwen38"]
        self.assertEqual((nodes.n, nodes.tp, nodes.fp, nodes.fn), (6, 3, 2, 1))
        self.assertEqual(nodes.reviewed_node_universe, 8)
        self.assertEqual(result.tables["corpus"].iloc[-1]["episodes"], 4)
        self.assertEqual(result.tables["corpus"].iloc[-1]["reviewed"], 3)
        self.assertEqual(result.tables["sensitivity"].query("cohort == 'tuning' and model_key == 'qwen38'").iloc[0]["n"], 1)

    def test_gate_uses_all_preceding_nodes_and_handles_node_four(self):
        result = build_paper_results(*fixture())
        rows = result.tables["sequential_node_agreement"].set_index("node_position")
        self.assertEqual(rows["eligible"].tolist(), [2, 2, 0, 0])
        for row in rows.itertuples():
            self.assertEqual(row.eligible, row.reviewer_complete + row.reviewer_incomplete)
            self.assertEqual(row.eligible, row.model_complete + row.model_incomplete)
        times = result.tables["timing_alignment"].set_index("model_key").loc["qwen38"]
        self.assertEqual((times.matched, times.within_2s, times.human_only, times.model_only), (3, 3, 1, 2))
        self.assertEqual(times.median_signed_error_s, 1)

    def test_empty_model_coverage_and_undefined_metrics_stay_explicit(self):
        episodes, nodes, observations = fixture()
        result = build_paper_results(episodes, nodes, observations.iloc[:0])
        self.assertTrue(result.tables["mission_metrics"].empty)
        self.assertTrue(result.tables["node_metrics"]["n"].eq(0).all())
        self.assertTrue(result.tables["node_metrics"]["balanced_accuracy"].isna().all())
        self.assertTrue(result.tables["sensitivity"]["model_episodes"].eq(0).all())
        figure = plot_timing(result)
        plt.close(figure)

    def test_publication_shapes_exports_plots_and_representative_evidence(self):
        result = build_paper_results(*fixture())
        tables = publication_tables(result)
        self.assertEqual(set(tables), {"corpus", "mission_metrics", "node_metrics", "sequential_node_agreement", "timing_alignment"})
        row = tables["node_metrics"].set_index("Auditor").loc["Qwen3.8"]
        self.assertEqual(row["Correct complete"], "3/4 (75.0%)")
        for level in ("mission", "node"):
            figure = plot_metric_by_scenario(result, level)
            plt.close(figure)
        figure = plot_timing(result)
        plt.close(figure)
        cases = representative_cases(result)
        self.assertEqual(len(cases), 2)
        self.assertTrue(cases["model_notes"].eq("evidence").all())
        with tempfile.TemporaryDirectory() as directory:
            paths = export_tables(result, directory, "03")
            self.assertTrue(all(p.exists() for p in paths))
            self.assertIn("\\toprule", (Path(directory) / "paper_mission_metrics.tex").read_text())
            exported = pd.read_csv(Path(directory) / "mission_metrics.csv")
            self.assertEqual(exported.iloc[0]["n"], 2)

    def test_duplicate_keys_rejected(self):
        episodes, nodes, obs = fixture()
        with self.assertRaises(ValueError):
            build_paper_results(pd.concat([episodes, episodes.iloc[:1]]), nodes, obs)
        with self.assertRaises(ValueError):
            build_paper_results(episodes, pd.concat([nodes, nodes.iloc[:1]]), obs)

    def test_endpoint_heatmap_uses_human_row_percentages(self):
        result = build_paper_results(*fixture())
        figure = plot_simulator_endpoint(result)
        self.assertEqual(len(figure.axes), 2)
        # Two human-complete episodes have simulator success, one incomplete does not.
        np.testing.assert_allclose(figure.axes[1].images[0].get_array(), [[100, 0], [0, 100]])
        self.assertIn("n=3", figure.axes[1].get_title())
        plt.close(figure)

    def test_prompt_figure_uses_common_full_corpus_but_reviewed_errors(self):
        episodes, nodes, obs = fixture()
        for frame in (episodes, nodes):
            frame.loc[frame["episode_dir"].eq("unreviewed"), "scenario"] = "Other task"
        inherited, tailored = obs.copy(), obs.copy()
        inherited["model_name"], tailored["model_name"] = "qwen35_qwen3_prompt", "qwen35"
        inherited = pd.concat([inherited, inherited[inherited["episode_dir"].eq("b")].assign(episode_dir="c")])
        tailored.loc[tailored["episode_dir"].eq("b"), "model_status_norm"] = "running"
        result = build_paper_results(episodes, nodes, pd.concat([obs, inherited, tailored]))
        rates = result.tables["prompt_completion_rates"].query("scenario == 'Overall'").set_index("model_key")
        self.assertEqual(rates.loc["qwen35_qwen3_prompt", "n"], 3)
        self.assertAlmostEqual(rates.loc["qwen35_qwen3_prompt", "complete_pct"], 200 / 3)
        self.assertAlmostEqual(rates.loc["qwen35", "complete_pct"], 100 / 3)
        errors = result.tables["prompt_error_tradeoff"].set_index("model_key")
        self.assertEqual(tuple(errors.loc["qwen35_qwen3_prompt", ["n", "fp", "fn"]]), (2, 1, 1))
        self.assertEqual(tuple(errors.loc["qwen35", ["n", "fp", "fn"]]), (2, 0, 1))
        figure = plot_prompt_sensitivity(result)
        self.assertEqual(len(figure.axes), 2)
        plt.close(figure)
        with self.assertRaisesRegex(ValueError, "No common"):
            plot_prompt_sensitivity(build_paper_results(*fixture()))

    def test_node_families_reuse_taxonomy_and_errors_use_all_pairs(self):
        episodes, nodes, obs = fixture()
        names = {
            "node_action_1": "node_action_grasp", "node_action_2": "node_action_place",
            "node_action_3": "node_action_close", "node_action_4": "node_action_handoff",
        }
        nodes["node_id"] = nodes["node_id"].map(names)
        obs["canonical_node_id"] = obs["canonical_node_id"].map(names)
        result = build_paper_results(episodes, nodes, obs)
        families = result.tables["node_family_metrics"].set_index("node_family")
        self.assertEqual(families.loc["Grasp and acquire", "n"], 2)
        self.assertEqual(families.loc["Transport, place, and handoff", "n"], 3)
        self.assertEqual(families.loc["State change or interaction", "n"], 1)
        self.assertEqual(families.loc["Grasp and acquire", "false_complete_pct_of_pairs"], 50)
        self.assertTrue(pd.isna(families.loc["State change or interaction", "balanced_accuracy"]))
        figure = plot_node_family_performance(result)
        self.assertEqual(figure.axes[0].images[0].get_array().shape, (3, 6))
        plt.close(figure)

    def test_temporal_event_union_excludes_both_incomplete_and_keeps_missing_clocks(self):
        episodes, nodes, obs = fixture()
        obs.loc[(obs["episode_dir"] == "b") & (obs["canonical_node_id"] == "node_action_2"), "model_status_norm"] = "running"
        obs.loc[(obs["episode_dir"] == "a") & (obs["canonical_node_id"] == "node_action_1"), "video_time_s"] = np.nan
        result = build_paper_results(episodes, nodes, obs)
        event = result.tables["completion_event_outcomes"].set_index("model_key").loc["qwen38"]
        self.assertEqual(event.event_union, 5)
        self.assertEqual(event.both_complete_missing_time, 1)
        self.assertEqual(event.event_union, event.within_2s + event.matched_outside_2s + event.human_only + event.model_only + event.both_complete_missing_time)
        figure = plot_timing(result)
        self.assertEqual(len(figure.axes), 2)
        heights = [patch.get_height() for patch in figure.axes[0].patches]
        self.assertAlmostEqual(sum(value for value in heights if np.isfinite(value)), 100)
        plt.close(figure)

    def test_sequential_bar_segments_reconcile_to_each_gate(self):
        result = build_paper_results(*fixture())
        counts = result.tables["sequential_completion_counts"]
        self.assertTrue(counts[list(dataset.nqcu.COMPLETION_CATEGORIES)].sum(axis=1).eq(counts.eligible_count).all())
        figure = plot_sequential_completion(result)
        self.assertEqual([axis.get_title() for axis in figure.axes], [
            "Node 1 Completion", "Node 2 Completion", "Node 3 Completion",
        ])
        plt.close(figure)

    def test_timing_boundaries_and_constant_or_single_samples(self):
        episodes, nodes, obs = fixture()
        completed = obs["episode_dir"].eq("a") & obs["model_status_norm"].eq("complete")
        obs.loc[completed, "video_time_s"] = [2, 2, 8.001]
        result = build_paper_results(episodes, nodes, obs)
        event = result.tables["completion_event_outcomes"].set_index("model_key").loc["qwen38"]
        self.assertEqual((event.within_2s, event.matched_outside_2s), (2, 1))
        for times, matched in (([2, 6, 8], 3), ([2, np.nan, np.nan], 1)):
            obs.loc[completed, "video_time_s"] = times
            result = build_paper_results(episodes, nodes, obs)
            event = result.tables["completion_event_outcomes"].set_index("model_key").loc["qwen38"]
            self.assertEqual(event.matched, matched)
            figure = plot_timing(result)
            self.assertEqual(len(figure.axes[1].collections), 1)  # One constant-value line, no KDE.
            plt.close(figure)

    def test_triptych_uses_csv_first_success_and_rebased_camera_frames(self):
        result = build_paper_results(*fixture())
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            camera = root / "frames" / "camera"
            camera.mkdir(parents=True)
            for step in range(24):
                Image.new("RGB", (8, 8), (step, 0, 0)).save(camera / f"frame_{step:05}.jpg")
            result.episodes.loc[result.episodes.episode_dir.eq("a"), [
                "episode_path", "episode_index", "run_tag", "episode_success_time",
                "episode_success_step", "episode_success_frame_path",
            ]] = [str(root), 3, "run", 9.2, 23, "/old/workstation/frames/camera/frame_00023.jpg"]
            figure, frames = plot_episode_triptych(result, "a")
            self.assertEqual(frames.frame_step.tolist(), [0, 11, 23])
            np.testing.assert_allclose(frames.machine_time_s, [0, 4.4, 9.2])
            self.assertTrue(frames.frame_path.map(lambda p: Path(p).is_file()).all())
            self.assertEqual(len(figure.axes), 3)
            plt.close(figure)
            (camera / "frame_00023.jpg").unlink()
            with self.assertRaises(FileNotFoundError):
                plot_episode_triptych(result, "a")

    def test_case_frames_are_spaced_and_image_gutters_match_border(self):
        with tempfile.TemporaryDirectory() as directory:
            camera = Path(directory) / "frames/camera"
            camera.mkdir(parents=True)
            for step in range(20):
                Image.new("RGB", (12, 8), (step, 0, 0)).save(camera / f"frame_{step:05}.jpg")
            for evidence, expected in (
                (12, [0, 7, 12, 19]), (1, [0, 1, 6, 19]), (19, [0, 14, 19]),
            ):
                with self.subTest(evidence=evidence):
                    case = pd.Series({
                        "episode_path": directory, "evidence_frame": f"frame_{evidence:05}",
                        "model": "Luna 6", "disagreement": "Missed complete",
                        "scenario": "Task", "node_id": "node_action",
                    })
                    figure = plot_case(case)
                    self.assertEqual(
                        [axis.get_title().splitlines()[0] for axis in figure.axes],
                        [f"frame_{step:05}.jpg" for step in expected],
                    )
                    boxes = [axis.get_position() for axis in figure.axes]
                    border = boxes[0].x0
                    self.assertAlmostEqual(border, 0.008)
                    self.assertAlmostEqual(1 - boxes[-1].x1, border)
                    for left, right in zip(boxes, boxes[1:]):
                        self.assertAlmostEqual(right.x0 - left.x1, border)
                    for axis in figure.axes:
                        box = axis.get_window_extent()
                        self.assertAlmostEqual(box.width / box.height, 1.5)
                    plt.close(figure)

    def test_progress_uses_complete_sheet0_and_only_valid_duration_overlap(self):
        episodes, nodes, obs = fixture()
        obs.loc[obs["episode_dir"].eq("a"), "video_time_s"] = [2, 3, 5, 7]
        result = build_paper_results(episodes, nodes, obs)
        points = pd.DataFrame({
            "t_sec": [2, 4], "progress_main": [0.5, 1.0], "progress_half": [0.25, 0.75],
        })

        result.episodes["reviewer_progress_points_json"] = points.to_json(orient="records")
        result.episodes["progress_review_total_rows"] = 2
        result.episodes["progress_review_scored_rows"] = [2, 1, 1, 1, 1]
        result.episodes["progress_review_duration_s"] = 4
        with patch("scripts.paper_results.review._read_sheet0_progress_points", side_effect=AssertionError("Raw workbook read")):
            coverage, detail = progress_alignment(result)
        self.assertEqual(int(coverage["progress_review_complete"].sum()), 1)
        self.assertEqual(set(detail["episode_dir"]), {"a"})
        self.assertEqual(detail["second"].tolist(), [2, 3, 4])
        self.assertEqual(detail["human_progress"].tolist(), [0.5, 0.75, 1.0])
        self.assertTrue(detail["model"].eq("Qwen3.8").all())

    def test_unmatched_and_missing_clock_events_are_not_zero_filled(self):
        episodes, nodes, obs = fixture()
        obs.loc[(obs["episode_dir"] == "a") & (obs["canonical_node_id"] == "node_action_1"), "video_time_s"] = float("nan")
        result = build_paper_results(episodes, nodes, obs)
        times = result.tables["timing_alignment"].set_index("model_key").loc["qwen38"]
        self.assertEqual((times.matched, times.both_complete_missing_time), (2, 1))
        self.assertEqual((times.human_only, times.model_only), (1, 2))

    def test_luna_timing_uses_measured_video_and_preserves_stored_clock(self):
        with tempfile.TemporaryDirectory() as directory:
            ep = Path(directory)
            (ep / "episode_summary.json").write_text(json.dumps({
                "steps_executed": 115, "outer_step_seconds": 0.4,
            }))
            (ep / "video").mkdir()
            (ep / "video/test.mp4").write_bytes(b"duration supplied by fixture")
            raw = pd.DataFrame({
                "node_id": ["node_action_grasp"] * 3,
                "machine_frame": ["frame_00000", "frame_00020", "frame_00114"],
                "video_time_m:ss": ["0:00", "0:08", "0:46"],
                "model_status": ["running", "complete", "complete"],
            })
            observations = dataset._standardize_observations(raw, "luna6_v5", dataset.LUNA6_SPEC, {})
            with patch.object(dataset, "_safe_video_duration_seconds", return_value=23.0):
                calibrated, clock = dataset._calibrate_luna_video_times(observations, ep)
                with self.assertRaisesRegex(ValueError, "frame index"):
                    dataset._calibrate_luna_video_times(
                        observations.assign(machine_frame="unverified"), ep,
                    )
            self.assertEqual(calibrated.video_time_s.tolist(), [0, 4, 23])
            self.assertEqual(calibrated.stored_video_time_s.tolist(), [0, 8, 46])
            self.assertEqual(observations.video_time_s.tolist(), [0, 8, 46])
            self.assertEqual(clock["simulation_to_video_scale"], 0.5)
            self.assertEqual(clock["video_seconds_per_saved_step"], 0.2)
            self.assertTrue(calibrated.timing_basis.eq("frame_index_scaled_to_video_duration").all())
            with patch.object(dataset, "_safe_video_duration_seconds", return_value=None):
                with self.assertWarnsRegex(RuntimeWarning, "Cannot measure video duration"):
                    unavailable, clock = dataset._calibrate_luna_video_times(observations, ep)
            self.assertTrue(unavailable.video_time_s.isna().all())
            self.assertEqual(unavailable.stored_video_time_s.tolist(), [0, 8, 46])
            self.assertEqual(unavailable["model_status_norm"].tolist(), ["running", "complete", "complete"])
            self.assertEqual(clock["status"], "unavailable")

    def test_loader_requires_completed_review_preserves_inputs_and_counts_conditions(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            name = "Task_episode_2_of_21_run"
            ep = root / name
            ep.mkdir()
            (ep / "episode_summary.json").write_text(json.dumps({
                "steps_executed": 10, "outer_step_seconds": 0.4,
            }))
            (ep / "bt.json").write_text(json.dumps({"root": {
                "type": "sequence", "children": [
                    {"id": "node_condition_ready", "type": "condition",
                     "condition": {"description": "Ready", "success_criteria": "Ready state"}},
                    {"id": "node_action_empty", "type": "action"},
                ],
            }}))
            (ep / "states_per_frame.jsonl").write_text('{"success":true}\n{"success":false}\n')
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1_BT_Node_Review"
            sheet.append(["node_order", "node_id", "reviewer_completion_time_m:ss"])
            sheet.append([1, "node_condition_ready", "0:00"])
            model = workbook.create_sheet("qwen38_v4")
            model.append(["node_id", "model_status", "video_time_m:ss"])
            model.append(["node_condition_ready", "complete", "0:00"])
            for sheet_name, status in (("luna6_v4", "running"), ("luna6_v5", "complete")):
                model = workbook.create_sheet(sheet_name)
                model.append(["node_id", "model_status", "video_time_m:ss", "machine_frame"])
                model.append(["node_condition_ready", status, "0:00", "frame_00000"])
            workbook.save(ep / "human_rater_evaluation.xlsx")
            workbook.close()
            (ep / "video").mkdir()
            (ep / "video/test.mp4").write_bytes(b"duration supplied by fixture")
            output = ep / "luna6_v5_test"
            output.mkdir()
            snapshot = b"frozen Luna v5"
            (output / "status_prompt_source.txt").write_bytes(snapshot)
            (output / "replay_config.json").write_text(json.dumps({
                "review_sheet_name": "luna6_v5",
                "status_prompt_sha256": hashlib.sha256(snapshot).hexdigest(),
                "status_max_new_tokens": 1024, "window_source": "frames",
                "expected_eval_every_n_steps": 5, "expected_status_window_frames": 10,
                "server_health": {
                    "backend": "cloud", "cloud_model": "gpt-6-luna",
                    "cloud_api": "responses", "cloud_reasoning_effort": "none",
                },
            }))
            schedule = root / "schedule.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "human_eval_schedule"
            sheet.append(["scenario_id", "run_tag", "episode_index", "episode_dir", "complete"])
            sheet.append(["Task", "run", 2, name, False])
            workbook.save(schedule)
            workbook.close()
            manifest = root / "manifest.txt"
            manifest.write_text(name + "\n")
            before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in root.rglob("*") if p.is_file()}
            with patch.object(dataset, "_safe_video_duration_seconds", return_value=2.0):
                episodes, observations = dataset.build_btaudit2_visualization_dataset(
                    logs_dir=root, schedule_path=schedule,
                )
            csv_path = root / "rows.csv"
            episodes.to_csv(csv_path, index=False)
            self.assertEqual(episodes.iloc[0]["qwen38_observation_count"], 1)
            self.assertEqual(len(json.loads(episodes.iloc[0]["qwen38_observations_json"])), 1)
            self.assertFalse(any(column.startswith("qwen38_v") for column in episodes))
            self.assertEqual(episodes.iloc[0]["qwen38_review_sheet"], "qwen38_v4")
            self.assertEqual(episodes.iloc[0]["luna6_review_sheet"], "luna6_v5")
            self.assertEqual(episodes.iloc[0]["luna6_observation_count"], 1)
            self.assertEqual(
                json.loads(episodes.iloc[0]["luna6_observations_json"])[0]["model_status_norm"],
                "complete",
            )
            embedded = json.loads(episodes.iloc[0]["luna6_observations_json"])[0]
            self.assertEqual(embedded["node_position"], 1)
            self.assertEqual(embedded["model_node_count"], 2)
            self.assertEqual(embedded["node_position_basis"], "saved_BT_ID_binding")
            self.assertEqual(
                json.loads(episodes.iloc[0]["model_node_contracts_json"])["luna6"]["nodes"][1]["node_id"],
                "node_action_empty",
            )
            self.assertEqual(embedded["stored_video_time_m:ss"], "0:00")
            self.assertEqual(embedded["stored_video_time_s"], 0)
            self.assertEqual(embedded["timing_basis"], "frame_index_scaled_to_video_duration")
            with patch("pandas.read_excel", side_effect=AssertionError("Notebook reopened workbook")):
                result = load_paper_results(csv_path, cohort_manifest=manifest)
            luna = result.node_model.query("model_key == 'luna6'")
            self.assertTrue(luna["model_complete"].all())
            self.assertEqual(luna["model"].unique().tolist(), ["Luna 6"])
            provenance = result.tables["replay_provenance"].query("model_key == 'luna6'").iloc[0]
            self.assertEqual(provenance.cloud_model, "gpt-6-luna")
            self.assertEqual(provenance.reasoning_effort, "none")
            self.assertEqual(provenance.window_frames, 10)
            self.assertTrue(provenance.snapshot_matches)
            self.assertEqual(provenance.analysis_time_basis, "frame_index_scaled_to_video_duration")
            self.assertEqual(provenance.analysis_video_seconds_per_step, 0.2)
            corpus = result.tables["corpus"].iloc[-1]
            self.assertEqual((corpus.episodes, corpus.reviewed, corpus.required_nodes), (1, 0, 1))
            self.assertEqual((corpus.success_ever, corpus.final_frame_success), (1, 0))
            self.assertTrue(result.tables["mission_metrics"].empty)
            self.assertEqual(result.nodes.iloc[0]["human_completion_s"], 0)
            self.assertEqual(result.node_model.iloc[0]["model_node_count"], 2)
            self.assertEqual(result.tables["auditability"].iloc[0]["structurally_incomplete_nodes"], 1)
            standardize = dataset._standardize_observations

            def blank_raw_ids(*args, **kwargs):
                return standardize(*args, **kwargs).assign(node_id="")

            with patch.object(dataset, "_standardize_observations", side_effect=blank_raw_ids), \
                    patch.object(dataset, "_safe_video_duration_seconds", return_value=2.0):
                blank_episodes, blank_observations = dataset.build_btaudit2_visualization_dataset(
                    logs_dir=root, schedule_path=schedule,
                )
            self.assertTrue(blank_observations.node_id.eq("").all())
            self.assertTrue(blank_observations.node_position.eq(1).all())
            self.assertEqual(
                json.loads(blank_episodes.iloc[0]["luna6_observations_json"])[0]["node_position"], 1,
            )
            after = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in before}
            self.assertEqual(before, after)
            manifest.write_text(name + "\n" + name + "\n")
            with self.assertRaises(ValueError):
                load_paper_results(csv_path, cohort_manifest=manifest)
            old = root / "old.csv"
            episodes.drop(columns="qwen38_observations_json").to_csv(old, index=False)
            with self.assertRaisesRegex(ValueError, "Refresh"):
                load_paper_results(old)


if __name__ == "__main__":
    unittest.main()
