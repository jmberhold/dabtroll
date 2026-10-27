"""Local release candidates preserve annotations, media, and immutable sources."""

import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import Workbook, load_workbook
import pandas as pd
from PIL import Image
import pyarrow.parquet as pq

from scripts import prepare_vibtarve_release as release


def fixture(root):
    root.mkdir()
    rows, observations = [], []
    for index in (1, 2):
        name = f"Task_Env_episode_{index}_of_21_run"
        episode = root / name
        episode.mkdir()
        record = {column: "[]" for column in release.CURRENT_SCHEMA_COLUMNS}
        record.update({
            "visualization_schema_version": release.SCHEMA_VERSION,
            "episode_dir": name, "episode_index": index, "scenario_id": "Task",
            "run_tag": "run", "schedule_complete_original": False,
            "high_level_task": "Move item", "bt_semantic_rating": None,
            "sim_success_episode": False, "sim_success_final": False,
            "required_nodes_json": json.dumps([{
                "node_id": "grasp", "description": "Grasp item", "success_criteria": "Item lifted",
                "required": True, "node_position": 1, "human_node_available": True,
                "human_completion_s": None, "human_complete": False, "reviewer_notes": "",
            }]),
            "luna6_time_calibration_json": "{}",
            "model_node_contracts_json": json.dumps({
                "luna6": {"node_count": 1, "bindings": {"grasp": 1},
                          "basis": "saved_BT_ID_binding"},
            }),
        })
        rows.append(record)
        for filename in ("bt.json", "bt_validation.json", "episode_manifest.json", "episode_outcome.json"):
            (episode / filename).write_text(json.dumps({"source": str(episode / filename)}))
        (episode / "episode_summary.json").write_text(json.dumps({
            "seed": 1, "steps_executed": 4, "outer_step_seconds": 0.4,
        }))
        (episode / "qwen_bt_raw.txt").write_text("Generated BT")
        for filename in ("states_per_frame.jsonl", "btstatus.jsonl"):
            (episode / filename).write_text('{"step":0,"success":false}\n')
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1_BT_Node_Review"
        sheet.append(["node_id", "reviewer_completion_time_m:ss", "reviewer_notes"])
        sheet.append(["grasp", None, "human note"])
        sheet["C2"].hyperlink = "https://example.org/evidence"
        sheet["C3"] = "=1+1"
        workbook.properties.creator = "Private Reviewer"
        workbook.properties.lastModifiedBy = "Private Reviewer"
        model = workbook.create_sheet("luna6_v2")
        model.append(["node_id", "model_status", "model_notes"])
        model.append(["grasp", "running", "historical model explanation"])
        workbook.save(episode / "human_rater_evaluation.xlsx")
        workbook.close()
        attempt = episode / "luna6_test"
        attempt.mkdir()
        prompt = b"Verify the visible grasp"
        (attempt / "status_prompt_source.txt").write_bytes(prompt)
        (attempt / "replay_config.json").write_text(json.dumps({
            "model": "luna", "review_sheet_name": "luna6_v5",
            "status_prompt_sha256": hashlib.sha256(prompt).hexdigest(),
            "server_health": {"host": "127.0.0.1"},
        }))
        for camera in ("ego", "cropped"):
            directory = episode / "frames" / camera
            directory.mkdir(parents=True)
            for step in range(4):
                Image.new("RGB", (8, 8), (step, 0, 0)).save(directory / f"frame_{step:05}.jpg")
        (episode / "video").mkdir()
        (episode / "video/test.mp4").write_bytes(b"probe supplied by test")
        observations.append({
            "episode_dir": name, "episode_index": index, "model_name": "luna6",
            "source": "mission_engine_response", "model_status_norm": "running",
            "video_time_s": 0, "stored_video_time_s": 0,
            "timing_basis": "frame_index_scaled_to_video_duration",
            "model_notes": "Item not lifted", "episode_path": str(episode),
        })
    pd.DataFrame(rows).to_csv(root / release.EPISODES_CSV, index=False)
    pd.DataFrame(observations).to_csv(root / release.OBSERVATIONS_CSV, index=False)
    pd.DataFrame([{"episode_dir": rows[0]["episode_dir"], "human_reviewed": True}]).to_csv(
        root / "btaudit2_molmo2_episode1_rows.csv", index=False)
    pd.DataFrame([{"code": "001", "path_json": json.dumps({"path": str(root / rows[0]["episode_dir"] / "bt.json")})}]).to_csv(
        root / "historical.csv", index=False)
    return rows


class ReleaseTests(unittest.TestCase):
    def test_episode_access_rejects_missing_reviews_unsafe_paths_and_corrupt_frames(self):
        with tempfile.TemporaryDirectory() as directory:
            root, candidate = Path(directory) / "btaudit2", Path(directory) / "candidate"
            fixture(root)
            with patch.object(release, "probe_video", return_value=2.0):
                release.prepare_release(root, candidate, expected_count=2, rights_confirmed=True)
            frame_path = candidate / "data/frames/evaluation.parquet"
            frames = pd.read_parquet(frame_path)
            for kind, message in (("path", "Unsafe"), ("hash", "checksum"), ("missing", "Sheet1")):
                with self.subTest(kind=kind):
                    if kind == "missing":
                        path = candidate / "data/workbook_rows/evaluation.parquet"
                        saved = path.read_bytes()
                        rows = pd.read_parquet(path)
                        rows[~rows.sheet.str.startswith("Sheet1")].to_parquet(path, index=False)
                    else:
                        path = frame_path
                        saved = path.read_bytes()
                        changed = frames.copy()
                        changed.loc[0, "member_path" if kind == "path" else "sha256"] = (
                            "../../escape.jpg" if kind == "path" else "0" * 64
                        )
                        changed.to_parquet(path, index=False)
                    try:
                        with self.assertRaisesRegex(ValueError, message):
                            release.stage_episode_access(candidate, Path(directory) / kind)
                    finally:
                        path.write_bytes(saved)

    def test_paper_bundle_is_portable_complete_and_source_preserving(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "bundle", Path(directory) / "release"
            source.mkdir()
            (source / "data").mkdir()
            csv = source / "data/episodes.csv"
            pd.DataFrame([{"code": "001", "path": "/home/operator/private.csv"}]).to_csv(csv, index=False)
            digest = release.sha256(csv)
            (source / "figure.pdf").write_bytes(b"%PDF-1.4\n%%EOF")
            (source / "manifest.json").write_text(json.dumps({
                "source": {"copy": "data/episodes.csv", "sha256": digest},
                "entries": [{"source_sha256": digest, "files": {"pdf": "figure.pdf"}}],
            }))
            (source / "index.html").write_text(f"<p>Full original episode CSV: {digest}</p>")
            original = {p: release.sha256(p) for p in source.rglob("*") if p.is_file()}
            info = release.stage_paper_bundle(source, output, release.Portable(Path(directory)))
            target = output / "paper"
            frame = pd.read_csv(target / "data/episodes.csv", dtype=str)
            self.assertEqual(frame.iloc[0]["code"], "001")
            self.assertNotIn("/home/", frame.iloc[0]["path"])
            catalog = json.loads((target / "manifest.json").read_text())
            self.assertEqual(catalog["source"]["original_sha256"], digest)
            self.assertEqual(catalog["source"]["sha256"], release.sha256(target / "data/episodes.csv"))
            self.assertEqual(catalog["entries"][0]["source_sha256"], catalog["source"]["sha256"])
            self.assertNotIn(digest, (target / "index.html").read_text())
            self.assertEqual((target / "figure.pdf").read_bytes(), (source / "figure.pdf").read_bytes())
            self.assertTrue(all(release.sha256(p) == value for p, value in original.items()))
            with zipfile.ZipFile(output / info["archive"]) as archive:
                self.assertIsNone(archive.testzip())
                for path in target.rglob("*"):
                    if path.is_file():
                        self.assertEqual(archive.read(path.relative_to(output).as_posix()), path.read_bytes())
            for line in (target / "SHA256SUMS").read_text().splitlines():
                value, name = line.split("  ", 1)
                self.assertEqual(release.sha256(target / name), value)
            with self.assertRaises(FileExistsError):
                release.stage_paper_bundle(source, output, release.Portable(Path(directory)))

    def test_paper_bundle_rejects_stale_sources_and_unsupported_assets(self):
        for unsupported in (False, True):
            with self.subTest(unsupported=unsupported), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "bundle"
                source.mkdir()
                (source / "data.csv").write_text("value\n1\n")
                (source / "manifest.json").write_text(json.dumps({
                    "source": {"copy": "data.csv", "sha256": "stale"}, "entries": [],
                }))
                if unsupported:
                    (source / "private.env").write_text("PRIVATE=yes")
                with self.assertRaisesRegex(ValueError, "Unsupported paper asset" if unsupported else "stale"):
                    release.stage_paper_bundle(source, Path(directory) / "release",
                                               release.Portable(Path(directory)))

    def test_local_candidate_preserves_human_review_and_all_summary_csvs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "btaudit2"
            fixture(root)
            original = {p: release.sha256(p) for p in root.rglob("*") if p.is_file()}
            output = Path(directory) / "candidate"
            with patch.object(release, "probe_video", side_effect=[2.0, None]):
                manifest = release.prepare_release(root, output, expected_count=2, rights_confirmed=True,
                                                   shard_bytes=4096)
            self.assertFalse((output / "INCOMPLETE").exists())
            self.assertEqual(manifest["frame_count"], 16)
            self.assertEqual(manifest["readable_videos"], 1)
            self.assertEqual(manifest["summary_csv_count"], 4)
            self.assertEqual(manifest["tables"]["human_annotations"]["rows"], 2)
            self.assertEqual(manifest["tables"]["workbook_rows"]["rows"], 10)
            self.assertFalse(manifest["upload_performed"])
            self.assertTrue(all(release.sha256(p) == digest for p, digest in original.items()))
            self.assertEqual(len(list(output.rglob("human_review.xlsx"))), 2)
            tuning = pd.read_parquet(output / "data/nodes/prompt_development.parquet")
            evaluation = pd.read_parquet(output / "data/nodes/evaluation.parquet")
            self.assertEqual(tuning.human_complete.tolist(), [False])
            self.assertTrue(evaluation.human_complete.isna().all())
            self.assertEqual(
                pq.read_schema(output / "data/nodes/prompt_development.parquet"),
                pq.read_schema(output / "data/nodes/evaluation.parquet"),
            )
            table = pd.read_csv(output / "summaries/historical.csv", dtype=str)
            self.assertEqual(table.iloc[0]["code"], "001")
            self.assertTrue(json.loads(table.iloc[0]["path_json"])["path"].startswith("episodes/"))
            attempts = pd.read_parquet(output / "data/replay_attempts/evaluation.parquet")
            episodes = pd.read_parquet(output / "data/episodes/evaluation.parquet")
            self.assertEqual(
                json.loads(episodes.iloc[0].model_node_contracts_json)["luna6"]["bindings"],
                {"grasp": 1},
            )
            self.assertEqual(json.loads(attempts.iloc[0].config_json)["server_health"]["host"],
                             "[PRIVATE_FIELD_REDACTED]")
            raw = pd.read_parquet(output / "data/workbook_rows/evaluation.parquet")
            self.assertIn("historical model explanation", raw.cells_json.str.cat())
            self.assertEqual(release.validate_release(output)["verified_frame_members"], 16)
            catalog = pd.read_csv(output / "episode_catalog.csv")
            self.assertEqual(catalog.human_reviewed.tolist(), [False, True])
            self.assertEqual(catalog.frame_count.tolist(), [8, 8])
            self.assertEqual(manifest["episode_access"]["completed_reviews"], {"prompt_development": 1})
            self.assertIn("Completed reviews (1)", (output / "EPISODES.md").read_text())
            for episode in catalog.to_dict("records"):
                book = load_workbook(output / episode["sheet1_xlsx"], data_only=False)
                self.assertEqual(book.sheetnames, ["Sheet1_BT_Node_Review"])
                self.assertEqual(book.active["C2"].value, "human note")
                self.assertEqual(book.active["C2"].hyperlink.target, "https://example.org/evidence")
                self.assertEqual(book.active["C3"].value, "=1+1")
                self.assertEqual(book.active["C3"].data_type, "s")
                self.assertEqual(book.properties.creator, "VIBTARVE")
                book.close()
                review = pd.read_csv(output / episode["sheet1_csv"])
                self.assertTrue(review.confirmed_review.eq(episode["human_reviewed"]).all())
                with zipfile.ZipFile(output / episode["frames_zip"]) as archive:
                    self.assertIsNone(archive.testzip())
                    self.assertEqual(len(archive.namelist()), 8)
                    for item in pd.read_csv(output / episode["frame_index"]).to_dict("records"):
                        self.assertEqual(hashlib.sha256(archive.read(item["zip_member_path"])).hexdigest(),
                                         item["sha256"])
            with self.assertRaises(FileExistsError):
                release.stage_episode_access(output, output)
            for line in (output / "SHA256SUMS").read_text().splitlines():
                digest, name = line.split("  ", 1)
                self.assertEqual(release.sha256(output / name), digest)
            with self.assertRaises(FileExistsError):
                release.prepare_release(root, output, expected_count=2, rights_confirmed=True)

    def test_privacy_is_recursive_and_suspicious_credentials_fail_explicitly(self):
        portable = release.Portable(Path("/tmp/btaudit2"))
        value = portable.clean({
            "path": "/home/operator/project/data/logs/btaudit2/Task_episode_1_of_21_run/bt.json",
            "email": "private@example.org", "endpoint": "http://localhost:5560",
            "notes": "See /home/operator/private/file.txt or http://127.0.0.1:8000",
        })
        self.assertEqual(value["path"], "episodes/Task_episode_1_of_21_run/bt.json")
        self.assertEqual(value["email"], "[PRIVATE_FIELD_REDACTED]")
        self.assertNotIn("/home/", value["notes"])
        self.assertNotIn("127.0.0.1", value["notes"])
        public_url = "https://huggingface.co/datasets/jmberhold/VIBTARVE"
        self.assertEqual(portable.clean(public_url), public_url)
        self.assertIsNone(release.LOCAL_PATH.search(public_url))
        self.assertEqual(portable.clean("file:///C:/Users/operator/private.txt"), "[LOCAL_PATH_REDACTED]")
        with self.assertRaisesRegex(ValueError, "Potential credential"):
            portable.clean("hf_" + "A" * 40)

    def test_unconfirmed_rights_and_source_output_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaisesRegex(ValueError, "confirmation"):
                release.prepare_release(root, root / "candidate")
            with self.assertRaisesRegex(ValueError, "outside"):
                release.prepare_release(root, root / "candidate", rights_confirmed=True)

    def test_card_does_not_advertise_empty_hugging_face_splits(self):
        info = {name: {"rows": 0, "cohorts": {}} for name in release.TABLES}
        info["episodes"] = {"rows": 2, "cohorts": {"prompt_development": 1, "evaluation": 1}}
        info["progress_annotations"] = {"rows": 1, "cohorts": {"evaluation": 1}}
        card = release.dataset_card({
            "repository": "jmberhold/VIBTARVE", "tables": info, "episode_count": 2,
            "cohorts": {"prompt_development": 1, "evaluation": 1}, "reviewed_evaluation_count": 1,
        })
        progress = card.split("- config_name: progress_annotations\n")[1].split("---", 1)[0]
        self.assertIn("split: evaluation", progress)
        self.assertNotIn("prompt_development.parquet", progress)
        self.assertNotIn("config_name: workbook_rows", card)
        self.assertIn("extra positions are not comparable, not false incomplete", card)
        self.assertIn("model_node_contracts_json", card)
        tags = card.split("tags:\n", 1)[1].split("configs:", 1)[0].splitlines()
        self.assertEqual(tags, [f"- {tag}" for tag in release.DATASET_TAGS])
        self.assertEqual(len(release.DATASET_TAGS), len(set(release.DATASET_TAGS)))
        self.assertTrue({"gr00t", "robocasa", "gr1-tabletop-tasks", "robot-simulation"}
                        .issubset(release.DATASET_TAGS))


if __name__ == "__main__":
    unittest.main()
