"""Prompt revisions and replay provenance without loading a model or simulation."""

import json
import os
import queue
import subprocess
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import numpy as np
import zmq
from openpyxl import Workbook, load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import prompts
import replay_qwen35_bt_eval as replay
from dabtroll_bt_pipeline import MissionEngineClient, PipelineConfig, _status_window_frame_count
import simulation
from dabtroll_bt_planner import build_status_messages


class PromptTests(unittest.TestCase):
    def test_luna_baseline_renders_user_template_without_changing_source(self):
        source = (ROOT / "config/status_prompts/luna6_v1.txt").read_text()
        profile = prompts.load_status_prompt("luna6_v1")
        self.assertEqual(profile.source_text, source.strip())
        rendered = profile("action", "Pick $cup", "Cup lifted")
        self.assertIn("Node description: Pick $cup", rendered)
        self.assertIn("Success criteria: Cup lifted", rendered)
        self.assertIn("Prior progress estimates, oldest to newest: none", rendered)
        self.assertNotIn("$node_type", rendered)
        self.assertEqual((ROOT / "config/status_prompts/luna6_v1.txt").read_text(), source)

    def test_baseline_is_exactly_current_standard_prompt(self):
        profile = prompts.load_status_prompt("qwen38_v1")
        for progress, criteria in ((None, None), ([0.1, 0.555], ["grasped", "lifted"])):
            args = ("action", "Pick {the} $cup", "Object lifted", progress, criteria)
            self.assertEqual(profile(*args), prompts.status_eval_text(*args))
        for name, builder in prompts.STATUS_PROMPT_BUILDERS.items():
            self.assertEqual(prompts.load_status_prompt(name)("a", "b", "c"), builder("a", "b", "c"))

    def test_template_validation_and_frozen_resolution(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "model_v2.txt"
            path.write_text("$node_type $node_desc $success_criteria")
            profile = prompts.load_status_prompt(str(path))
            path.write_text("changed $node_type $node_desc $success_criteria")
            self.assertEqual(profile("a", "b", "c"), "a b c")
            self.assertNotEqual(profile.sha256, prompts.load_status_prompt(str(path)).sha256)
            for text in ("$unknown $node_type $node_desc $success_criteria", "$node_type", "$ node_type"):
                path.write_text(text)
                with self.assertRaises(ValueError):
                    prompts.load_status_prompt(str(path))
        with self.assertRaises(ValueError):
            prompts.load_status_prompt("does_not_exist")

    def test_direct_planner_uses_same_prompt_selector(self):
        messages = build_status_messages(["frame.jpg"], "action", "pick", "lift", status_prompt="qwen38_v1")
        self.assertEqual(messages[0]["content"][0], {"type": "image", "image": "frame.jpg"})
        self.assertEqual(messages[0]["content"][-1]["text"], prompts.status_eval_text("action", "pick", "lift"))

    def test_live_client_uses_selected_prompt_and_history(self):
        client = MissionEngineClient(status_prompt="qwen38_v1")
        try:
            response = {"ok": True, "text": json.dumps({"status": "running", "progress_score": 0.7})}
            with patch.object(client, "request", return_value=response) as request:
                result, trace = client.evaluate_node_status_with_trace(
                    {"type": "action", "action": {"description": "pick", "success_criteria": "lift"}},
                    [np.zeros((16, 16, 3), dtype=np.uint8)], [0.6], ["near cup"],
                )
            self.assertEqual(
                request.call_args.args[0]["prompt_text"],
                prompts.status_eval_text("action", "pick", "lift", [0.6], ["near cup"]),
            )
            self.assertEqual(result["_vlm_prompt_id"], "qwen38_v1")
            self.assertEqual(trace[0]["status_prompt"]["sha256"], client.status_prompt.sha256)
        finally:
            client.sock.close(linger=0)


class ReplayTests(unittest.TestCase):
    def test_seconds_overrides_replace_recorded_counts(self):
        summary = {
            "outer_step_seconds": 0.4, "status_eval_every_n_steps": 5,
            "status_eval_seconds_target": 2.0, "status_window_frames": 8,
            "status_window_seconds_target": 4.0,
        }
        for overrides, expected in (
            ((None, None, None), (5, 8)),
            ((2.0, 4.0, None), (5, 10)),
            ((1.2, 6.0, None), (3, 15)),
            ((None, 4.0, 8), (5, 8)),
            ((None, 4.0, 0), (5, 10)),
        ):
            with self.subTest(overrides=overrides):
                self.assertEqual(replay._derive_timing_from_summary(summary, *overrides), expected)
        self.assertEqual(
            replay._derive_timing_from_summary({"outer_step_seconds": 0.3}, 1.2, 3.0, None),
            (4, 10),
        )

    def test_four_second_windows_contain_ten_ordered_samples(self):
        frames = self.episode / "frames" / "camera"
        for index in range(3, 16):
            Image.new("RGB", (16, 16), "red").save(frames / f"frame_{index:05d}.jpg")
        rows = replay._build_manifest_rows_from_frames(self.episode, "test", 5, 10, 14)
        self.assertEqual([row["step"] for row in rows], [5, 10, 14])
        self.assertEqual([len(row["frames"]) for row in rows], [6, 10, 10])
        self.assertEqual(
            [Path(p).name for p in rows[1]["frames"]],
            [f"frame_{index:05d}.jpg" for index in range(1, 11)],
        )
        self.assertEqual(
            [Path(p).name for p in rows[-1]["frames"]],
            [f"frame_{index:05d}.jpg" for index in range(5, 15)],
        )
        (frames / "frame_00007.jpg").unlink()
        with self.assertRaisesRegex(ValueError, "Incomplete replay frame window"):
            replay._build_manifest_rows_from_frames(self.episode, "test", 5, 10, 14)

    def test_invalid_seconds_overrides_are_rejected(self):
        for value in (0, -1, float("inf"), float("nan")):
            for overrides in ((value, None, None), (None, value, None)):
                with self.subTest(overrides=overrides):
                    with self.assertRaisesRegex(ValueError, "finite and positive"):
                        replay._derive_timing_from_summary({}, *overrides)

    def test_manifest_timing_overrides_fail_before_contacting_server(self):
        (self.episode / "status_window_manifest.jsonl").write_text("")
        with patch.object(replay, "MissionEngineClient") as client:
            with self.assertRaisesRegex(ValueError, "--window-source frames"):
                replay._replay_episode(**{**self.args, "window_source": "manifest"})
        client.assert_not_called()

    def test_historical_manifest_replay_preserves_recorded_evidence(self):
        rows = replay._build_manifest_rows_from_frames(self.episode, "test", 1, 2, 2)
        (self.episode / "status_window_manifest.jsonl").write_text(
            "".join(json.dumps(row) + "\n" for row in rows)
        )
        client = Mock()
        client.request.side_effect = [self.health, self.response]
        with patch.object(replay, "MissionEngineClient", return_value=client):
            result = replay._replay_episode(**{
                **self.args, "window_source": "manifest",
                "status_eval_seconds_override": None,
                "status_window_frames_override": None,
            })
        output = Path(result["replay_output_dir"])
        saved = json.loads((output / "status_window_manifest.jsonl").read_text().splitlines()[0])
        self.assertEqual(saved["frames"], rows[0]["frames"])
        self.assertEqual(saved["frame_timing"]["frame_count"], 2)

    def test_simulation_preserves_seconds_derived_window(self):
        for extra, expected in (([], 10), (["--status-window-frames", "8"], 8)):
            with self.subTest(extra=extra):
                argv = [
                    "simulation.py", "--env-name", "test", "--test", "btaudit2",
                    "--project-root", str(self.root),
                    "--status-eval-seconds", "2", "--status-window-seconds", "4", *extra,
                ]
                with (
                    patch.object(sys, "argv", argv),
                    patch.object(simulation, "run_dabtroll_episode", return_value={
                        "bt_json_path": str(self.episode / "bt.json"),
                    }) as run,
                    patch.object(simulation, "_auto_generate_eval_workbook"),
                    patch("builtins.print"),
                ):
                    simulation.main()
                cfg = run.call_args.args[0]
                self.assertEqual(_status_window_frame_count(cfg), expected)
        cfg = PipelineConfig(env_name="test", n_action_steps=8, status_window_seconds=4)
        self.assertEqual(_status_window_frame_count(cfg), 10)

    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.episode = self.root / "task_episode_1_of_21_test"
        self.episode.mkdir()
        (self.episode / "episode_summary.json").write_text(json.dumps({
            "episode_id": "test", "mission_name": "test mission", "steps_executed": 3,
            "outer_step_seconds": 0.4, "n_action_steps": 8, "control_freq_hz": 20,
            "status_eval_seconds_target": 0.4, "status_window_frames": 2,
        }))
        (self.episode / "bt.json").write_text(json.dumps({
            "root": {"id": "root", "type": "sequence", "children": [{
                "id": "pick", "type": "action", "description": "Pick",
                "action": {"id": "pick", "description": "Pick cup", "success_criteria": "Cup lifted"},
                "children": [],
            }]},
        }))
        frames = self.episode / "frames" / "camera"
        frames.mkdir(parents=True)
        for index in range(3):
            Image.new("RGB", (16, 16), "red").save(frames / f"frame_{index:05d}.jpg")
        wb = Workbook()
        wb.active.title = "human_ratings"
        wb.active["A1"] = "Keep this human review"
        wb.save(self.episode / "human_rater_evaluation.xlsx")
        self.health = {"ok": True, "model": "qwen38", "video_input": "contact_sheet"}
        self.response = {"ok": True, "text": json.dumps({
            "status": "complete", "progress_score": 1.0, "notes": "visible",
            "criteria_met": ["Cup lifted"], "criteria_missing": [],
        })}
        self.args = dict(
            episode_dir=self.episode, mission_host="127.0.0.1", mission_port=5560,
            mission_timeout_ms=5000, out_tag="test", status_eval_seconds_override=0.4,
            status_window_seconds_override=None, status_window_frames_override=2,
            use_manifest_prompts=False, window_source="frames", review_sheet_name="qwen38_v1",
            status_prompt="qwen38_v1", model_label="quantized Qwen3.8", output_prefix="qwen38_v1",
            status_max_new_tokens=256, expected_mission_model="qwen38", overwrite_review_sheet=False,
        )

    def test_replay_metadata_workbook_and_overwrite_guard(self):
        client = Mock()
        client.request.side_effect = [self.health, self.response]
        with patch.object(replay, "MissionEngineClient", return_value=client):
            result = replay._replay_episode(**self.args)
        self.assertFalse(result["mission_has_errors"])
        output = Path(result["replay_output_dir"])
        config = json.loads((output / "replay_config.json").read_text())
        self.assertEqual(config["status_prompt"], "qwen38_v1")
        self.assertEqual(config["server_health"], self.health)
        self.assertEqual(config["timing"]["clock"], "simulation")
        self.assertEqual(config["timing"]["resolved_window_nominal_seconds"], 0.8)
        self.assertEqual(config["timing"]["resolved_window_first_to_last_seconds"], 0.4)
        rows = [json.loads(line) for line in (output / "missionengine.jsonl").read_text().splitlines()]
        self.assertTrue(all(row["status_prompt_sha256"] == prompts.load_status_prompt("qwen38_v1").sha256 for row in rows))
        request = next(row for row in rows if row["direction"] == "request")
        self.assertEqual(request["frame_timing"]["frame_count"], 2)
        self.assertAlmostEqual(request["frame_timing"]["first_to_last_seconds"], 0.4)
        workbook = load_workbook(self.episode / "human_rater_evaluation.xlsx")
        try:
            self.assertEqual(workbook["human_ratings"]["A1"].value, "Keep this human review")
            self.assertIn("quantized Qwen3.8", workbook["qwen38_v1"]["A6"].value)
            self.assertIn("qwen38_v1", workbook["qwen38_v1"]["A6"].value)
        finally:
            workbook.close()
        with self.assertRaises(FileExistsError):
            replay._replay_episode(**self.args)
        with self.assertRaisesRegex(ValueError, "changed since"):
            replay._replay_episode(**{**self.args, "expected_prompt_sha256": "wrong"})

    def test_invalid_status_and_wrong_model_are_not_clean(self):
        client = Mock()
        client.request.side_effect = lambda request, **kwargs: (
            self.health if request["mode"] == "health" else {"ok": True, "text": "not status JSON"}
        )
        with patch.object(replay, "MissionEngineClient", return_value=client):
            result = replay._replay_episode(**self.args)
        self.assertTrue(result["mission_has_errors"])
        self.assertFalse(result["review_sheet_written"])
        self.assertGreater(result["mission_invalid_status_rows"], 0)
        client.request.side_effect = [{"ok": True, "model": "other"}]
        with patch.object(replay, "MissionEngineClient", return_value=client):
            with self.assertRaisesRegex(ValueError, "Expected mission model"):
                replay._replay_episode(**{**self.args, "out_tag": "other"})

    def test_engine_failure_stops_episode_and_preserves_error_artifacts(self):
        client = Mock()
        client.request.side_effect = [
            self.health,
            {"ok": False, "error": "EngineCore encountered an issue. See stack trace."},
        ]
        with patch.object(replay, "MissionEngineClient", return_value=client):
            result = replay._replay_episode(**{**self.args, "stop_on_backend_failure": True})
        self.assertEqual(client.request.call_count, 2)
        self.assertEqual(result["mission_backend_error_rows"], 1)
        self.assertFalse(result["review_sheet_written"])
        self.assertTrue(Path(result["replay_output_dir"], "missionengine.jsonl").exists())

    def test_unreadable_saved_video_fails_before_inference(self):
        video = self.episode / "video"
        video.mkdir()
        (video / "test.mp4").write_bytes(b"unreadable video")
        with patch.object(replay, "_safe_video_duration_seconds", return_value=None), \
                patch.object(replay, "MissionEngineClient") as client:
            with self.assertRaisesRegex(ValueError, "simulation seconds must not be labeled video time"):
                replay._replay_episode(**self.args)
        client.assert_not_called()

    def test_shell_experiment_freezes_prompt_and_writes_valid_json(self):
        self._assert_shell_experiment()

    def test_shell_stops_on_dead_backend_and_keeps_remaining_episodes(self):
        self._assert_shell_experiment(backend_failure=True)

    def test_luna_shell_freezes_prompt_and_records_cloud_provenance(self):
        self._assert_shell_experiment(model="luna")

    def test_luna_shell_stops_on_provider_error_without_retrying(self):
        self._assert_shell_experiment(backend_failure=True, model="luna")

    def test_cloud_backend_errors_are_recognized_for_fail_fast(self):
        for message in (
            "Error code: 401 - invalid API key",
            "Error code: 429 - rate limit exceeded",
            "Error code: 500 - provider unavailable",
            "Request timed out.",
            "Cloud response was not completed (max_output_tokens); do not use partial output.",
            "Cloud model returned no usage; inspect provider billing before retrying.",
        ):
            with self.subTest(message=message):
                self.assertIsNotNone(replay.BACKEND_ERROR_RE.search(message))
        self.assertIsNone(replay.BACKEND_ERROR_RE.search("Object has not reached the cabinet."))

    def _assert_shell_experiment(self, backend_failure=False, model="qwen38"):
        prompt_id = "luna6_v1" if model == "luna" else "qwen38_v1"
        launcher = ROOT / f"scripts/run_replay_btaudit2_{model}_episode1.sh"
        self.health["model"] = model
        if model == "luna":
            self.health.update(
                backend="cloud", cloud_model="gpt-6-luna", cloud_reasoning_effort="none",
                video_input="frames",
            )
            self.response.update(backend="cloud", cloud_model="gpt-6-luna", usage={"input_tokens": 90})
        backend_error = "Error code: 429 - rate limit exceeded" if model == "luna" else "EngineCore encountered an issue"
        addresses = queue.Queue()
        requests = queue.Queue()
        stop = threading.Event()

        def serve():
            context = zmq.Context()
            sock = context.socket(zmq.REP)
            sock.LINGER = 0
            port = sock.bind_to_random_port("tcp://127.0.0.1")
            addresses.put(port)
            try:
                while not stop.is_set():
                    if sock.poll(100):
                        request = sock.recv_json()
                        requests.put(request)
                        response = (
                            {"ok": False, "error": backend_error}
                            if backend_failure else self.response
                        )
                        sock.send_json(self.health if request["mode"] == "health" else response)
            finally:
                sock.close()
                context.term()

        thread = threading.Thread(target=serve)
        thread.start()
        port = addresses.get(timeout=5)
        if backend_failure:
            import shutil

            shutil.copytree(self.episode, self.root / "other_episode_1_of_21_test")
        env = {
            **os.environ,
            "PYTHON_BIN": sys.executable,
            "REPLAY_SCRIPT": str(ROOT / "scripts/replay_qwen35_bt_eval.py"),
            "LOGS_ROOT": str(self.root),
            "STATUS_PROMPT": prompt_id, "REVIEW_SHEET_NAME": prompt_id,
            "MISSION_HOST": "127.0.0.1", "MISSION_PORT": str(port),
            "EPISODE_MANIFEST": "", "MAX_PASSES": "1",
        }
        for key in ("STATUS_EVAL_SECONDS", "STATUS_WINDOW_SECONDS", "STATUS_WINDOW_FRAMES"):
            env.pop(key, None)
        self.response = {"ok": True, "text": json.dumps({"status": "running", "progress_score": 0.1})}
        summary_path = self.episode / "episode_summary.json"
        summary = json.loads(summary_path.read_text())
        summary.update(
            steps_executed=15, status_eval_every_n_steps=5,
            status_eval_seconds_target=2.0, status_window_frames=8,
            status_window_seconds_target=4.0,
        )
        summary_path.write_text(json.dumps(summary))
        for ep in self.root.glob("*episode_1_of_21_*"):
            (ep / "episode_summary.json").write_text(json.dumps(summary))
            for index in range(3, 15):
                Image.new("RGB", (16, 16), "red").save(ep / "frames/camera" / f"frame_{index:05d}.jpg")
        env.pop("STATUS_MAX_NEW_TOKENS", None)
        noisy_replay = self.root / "noisy_replay.py"
        noisy_replay.write_text(
            "import runpy, sys\n"
            f"sys.path.insert(0, {str(ROOT / 'scripts')!r})\n"
            "print('[diagnostic] replay starting')\n"
            f"runpy.run_path({str(ROOT / 'scripts/replay_qwen35_bt_eval.py')!r}, run_name='__main__')\n"
        )
        env["REPLAY_SCRIPT"] = str(noisy_replay)
        env["PYTHONPATH"] = str(ROOT / "scripts") + os.pathsep + env.get("PYTHONPATH", "")
        try:
            run = subprocess.run(
                ["bash", str(launcher), "baseline_test"],
                env=env, capture_output=True, text=True, timeout=120,
            )
        finally:
            stop.set()
            thread.join(timeout=5)
        self.assertEqual(run.returncode, 1 if backend_failure else 0, run.stdout + run.stderr)
        run_root = self.root / "replay_episode1_mod_runs" / "baseline_test"
        metadata = json.loads((run_root / "experiment.json").read_text())
        self.assertEqual(metadata["settings"]["STATUS_MAX_NEW_TOKENS"], "1024")
        self.assertEqual(metadata["settings"]["STATUS_EVAL_SECONDS"], "2")
        self.assertEqual(metadata["settings"]["STATUS_WINDOW_SECONDS"], "4")
        status_requests = []
        while not requests.empty():
            request = requests.get_nowait()
            if request["mode"] == "text_video":
                status_requests.append(request)
        self.assertTrue(status_requests)
        self.assertTrue(all(request["max_new_tokens"] == 1024 for request in status_requests))
        if not backend_failure:
            self.assertEqual([len(request["frames_b64"]) for request in status_requests], [6, 10, 10])
        self.assertEqual(metadata["prompt"]["sha256"], prompts.load_status_prompt(prompt_id).sha256)
        self.assertEqual(metadata["server_health"]["model"], model)
        if model == "luna":
            self.assertEqual(metadata["settings"]["MODEL_LABEL"], "gpt-6-luna")
            self.assertEqual(metadata["settings"]["MISSION_TIMEOUT_MS"], "120000")
            self.assertEqual(metadata["server_health"]["cloud_model"], "gpt-6-luna")
        snapshot = prompts.load_status_prompt(metadata["prompt_snapshot"])
        self.assertEqual(snapshot.sha256, metadata["prompt"]["sha256"])
        detail = json.loads((run_root / "pass_01/details.jsonl").read_text())
        config = json.loads((Path(detail["qwen_output_dir"]) / "replay_config.json").read_text())
        self.assertEqual(config["expected_eval_every_n_steps"], 5)
        self.assertEqual(config["expected_status_window_frames"], 10)
        self.assertEqual(config["timing"]["resolved_eval_seconds"], 2)
        self.assertEqual(config["timing"]["resolved_window_nominal_seconds"], 4)
        self.assertEqual(config["timing"]["resolved_window_first_to_last_seconds"], 3.6)
        self.assertEqual(detail["review_sheet_written"], not backend_failure)
        self.assertIn(
            f"remaining_count={2 if backend_failure else 0}",
            (run_root / "run_summary.txt").read_text(),
        )
        if backend_failure:
            self.assertEqual(detail["mission_backend_error_rows"], 1)
            self.assertEqual(len((run_root / "pending_pass_01.txt").read_text().splitlines()), 2)
        duplicate = subprocess.run(
            ["bash", str(launcher), "baseline_test"],
            env=env, capture_output=True, text=True, timeout=10,
        )
        self.assertNotEqual(duplicate.returncode, 0)


if __name__ == "__main__":
    unittest.main()
