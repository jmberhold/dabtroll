"""Read-only structural and replay-provenance audit with portable cohort manifests."""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook

from replay_qwen35_bt_eval import _summarize_mission_response_quality


EPISODE_RE = re.compile(r"(?P<scenario>.+)_episode_(?P<index>\d+)_of_(?P<count>\d+)_(?P<run>.+)")
REQUIRED_FILES = (
    "episode_summary.json", "episode_manifest.json", "episode_outcome.json",
    "bt.json", "bt_validation.json", "start_frame.jpg", "bt.svg",
    "pipeline_trace.jsonl", "btstatus.jsonl", "states_per_frame.jsonl",
    "missionengine.jsonl", "status_window_manifest.jsonl", "human_rater_evaluation.xlsx",
)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = []
    for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        row = json.loads(line)
        if not isinstance(row, dict):
            raise ValueError(f"{path.name}:{number}: expected a JSON object")
        rows.append(row)
    return rows


def audit_episode(
    episode: Path, root: Path, review_sheet: str = "",
    prompt_sha256: str = "", max_new_tokens: int | None = None,
) -> dict[str, Any]:
    match = EPISODE_RE.fullmatch(episode.name)
    if match is None:
        raise ValueError(f"Not an episode directory: {episode}")
    row: dict[str, Any] = {
        "episode": episode.relative_to(root).as_posix(),
        "scenario": match["scenario"], "episode_index": int(match["index"]),
        "run_tag": match["run"], "issues": [], "warnings": [],
        "metadata_sha256": {}, "replay_attempts": [],
    }
    issues = row["issues"]
    for name in REQUIRED_FILES:
        path = episode / name
        if not path.is_file():
            issues.append(f"missing:{name}")
            continue
        row["metadata_sha256"][name] = file_sha256(path)
        if path.suffix == ".json":
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                if not isinstance(data, dict):
                    raise ValueError("expected a JSON object")
                if name == "episode_summary.json":
                    row["episode_id"] = data.get("episode_id")
                    row["seed"] = data.get("seed")
                    if data.get("episode_index") != int(match["index"]):
                        issues.append("summary_episode_index_mismatch")
                if name == "bt.json" and not isinstance(data.get("root"), dict):
                    issues.append("bt_missing_root")
            except (ValueError, OSError, BadZipFile) as exc:
                issues.append(f"invalid:{name}:{exc}")
    row["frame_count"] = sum(
        1 for path in (episode / "frames").rglob("*")
        if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".png"}
    )
    if not row["frame_count"]:
        issues.append("missing_frames")
    if any(episode.glob(".~lock.*")) or any(episode.glob("~$*")):
        row["warnings"].append("workbook_lock_present: do not modify while a reviewer has it open")

    workbook = episode / "human_rater_evaluation.xlsx"
    if workbook.is_file():
        wb = load_workbook(workbook, read_only=True, data_only=False)
        try:
            row["sheets"] = wb.sheetnames
            row["human_sheet_fingerprints"] = {
                name: hashlib.sha256(json.dumps(list(wb[name].values), default=str).encode()).hexdigest()
                for name in wb.sheetnames if name.lower().startswith(("sheet0", "sheet1"))
            }
            if not any(name.lower().startswith("sheet1") for name in wb.sheetnames):
                issues.append("missing_human_node_sheet")
            if review_sheet:
                if review_sheet not in wb.sheetnames:
                    issues.append(f"missing_review_sheet:{review_sheet}")
                elif prompt_sha256 and not any(
                    prompt_sha256 in str(value)
                    for values in wb[review_sheet].iter_rows(max_row=8, values_only=True)
                    for value in values
                ):
                    issues.append("review_sheet_prompt_hash_mismatch")
        finally:
            wb.close()

    if review_sheet:
        for config_path in sorted(episode.glob("*/replay_config.json")):
            try:
                config = json.loads(config_path.read_text(encoding="utf-8"))
                if config.get("review_sheet_name") != review_sheet:
                    continue
                attempt: dict[str, Any] = {"path": config_path.parent.relative_to(root).as_posix()}
                attempt["prompt_sha256"] = config.get("status_prompt_sha256")
                attempt["max_new_tokens"] = config.get("status_max_new_tokens")
                attempt["model"] = config.get("model")
                quality = _summarize_mission_response_quality(read_jsonl(config_path.parent / "missionengine.jsonl"))
                attempt["quality"] = quality
                attempt["matches_settings"] = (
                    (not prompt_sha256 or attempt["prompt_sha256"] == prompt_sha256)
                    and (max_new_tokens is None or attempt["max_new_tokens"] == max_new_tokens)
                )
                snapshot = config_path.parent / "status_prompt_source.txt"
                attempt["snapshot_matches"] = snapshot.is_file() and file_sha256(snapshot) == attempt["prompt_sha256"]
                row["replay_attempts"].append(attempt)
            except (ValueError, OSError) as exc:
                issues.append(f"invalid_replay:{config_path.parent.name}:{exc}")
        clean = [
            attempt for attempt in row["replay_attempts"]
            if attempt["matches_settings"] and attempt["snapshot_matches"]
            and not attempt["quality"]["mission_has_errors"]
        ]
        row["matching_clean_attempts"] = len(clean)
        if len(clean) != 1:
            issues.append(f"expected_one_matching_clean_replay:found={len(clean)}")
        else:
            row["clean_output"] = clean[0]["path"]
            row["clean_responses"] = clean[0]["quality"]["mission_response_rows"]
    return row


def audit_dataset(
    root: Path, episode_min: int = 1, episode_max: int = 21,
    review_sheet: str = "", prompt_sha256: str = "", max_new_tokens: int | None = None,
    manifest: Path | None = None,
) -> dict[str, Any]:
    root = root.expanduser().resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"Dataset root not found: {root}")
    if not 1 <= episode_min <= episode_max:
        raise ValueError("Require 1 <= episode-min <= episode-max")
    if manifest:
        entries = [line.strip() for line in manifest.read_text().splitlines() if line.strip()]
        paths = []
        for entry in entries:
            relative = Path(entry)
            if relative.is_absolute() or ".." in relative.parts:
                raise ValueError(f"Portable manifests require dataset-relative paths: {entry}")
            path = (root / relative).resolve()
            if not path.is_relative_to(root) or path.parent != root or not path.is_dir():
                raise ValueError(f"Manifest episode is missing or outside dataset root: {entry}")
            paths.append(path)
        if len(paths) != len(set(paths)):
            raise ValueError("Duplicate episodes in manifest")
    else:
        paths = [path for path in root.iterdir() if path.is_dir() and EPISODE_RE.fullmatch(path.name)]
    selected = [
        path for path in paths
        if (match := EPISODE_RE.fullmatch(path.name)) and episode_min <= int(match["index"]) <= episode_max
    ]
    if not selected:
        raise ValueError("No episodes matched the requested cohort")
    rows = []
    for episode in sorted(selected):
        try:
            rows.append(audit_episode(episode, root, review_sheet, prompt_sha256, max_new_tokens))
        except (ValueError, OSError) as exc:
            rows.append({"episode": episode.name, "issues": [f"unreadable:{exc}"], "warnings": []})
    return {
        "schema": "dabtroll_dataset_audit_v1", "episode_count": len(rows),
        "issue_episode_count": sum(bool(row["issues"]) for row in rows),
        "clean_responses": sum(row.get("clean_responses", 0) for row in rows),
        "scenario_counts": dict(Counter(row.get("scenario", "unreadable") for row in rows)),
        "review_sheet": review_sheet, "prompt_sha256": prompt_sha256,
        "max_new_tokens": max_new_tokens, "episodes": rows,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset-root", type=Path, default=Path(__file__).resolve().parents[1] / "data/logs/btaudit2")
    parser.add_argument("--output-dir", required=True, type=Path, help="New report directory; never overwritten.")
    parser.add_argument("--manifest", type=Path, help="Optional portable, dataset-relative cohort manifest.")
    parser.add_argument("--episode-min", type=int, default=1)
    parser.add_argument("--episode-max", type=int, default=21)
    parser.add_argument("--review-sheet", default="")
    parser.add_argument("--prompt-sha256", default="")
    parser.add_argument("--max-new-tokens", type=int)
    parser.add_argument("--expected-count", type=int)
    args = parser.parse_args()
    if (args.prompt_sha256 or args.max_new_tokens is not None) and not args.review_sheet:
        parser.error("--prompt-sha256/--max-new-tokens require --review-sheet")
    if args.max_new_tokens is not None and args.max_new_tokens <= 0:
        parser.error("--max-new-tokens must be positive")
    if args.prompt_sha256 and not re.fullmatch(r"[0-9a-f]{64}", args.prompt_sha256):
        parser.error("--prompt-sha256 must be a lowercase SHA256 digest")
    if args.output_dir.exists():
        parser.error("Output directory already exists; choose a new audit run directory")
    report = audit_dataset(
        args.dataset_root, args.episode_min, args.episode_max, args.review_sheet,
        args.prompt_sha256, args.max_new_tokens, args.manifest,
    )
    report["count_matches_expected"] = args.expected_count is None or report["episode_count"] == args.expected_count
    args.output_dir.mkdir(parents=True)
    (args.output_dir / "audit.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    names = [row["episode"] for row in report["episodes"]]
    (args.output_dir / "episodes_relative.txt").write_text("".join(name + "\n" for name in names), encoding="utf-8")
    root = args.dataset_root.expanduser().resolve()
    (args.output_dir / "episodes_local.txt").write_text("".join(str(root / name) + "\n" for name in names), encoding="utf-8")
    (args.output_dir / "pending_local.txt").write_text(
        "".join(str(root / row["episode"]) + "\n" for row in report["episodes"] if row["issues"]),
        encoding="utf-8",
    )
    print(json.dumps({key: value for key, value in report.items() if key != "episodes"}, indent=2))
    if report["issue_episode_count"] or not report["count_matches_expected"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
