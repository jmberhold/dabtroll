#!/usr/bin/env python3
"""Audit legacy Sheet2 and status-response quality without modifying episode files."""

from __future__ import annotations

import argparse
import json
import re
import time
import warnings
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook


EPISODE_RE = re.compile(r"_episode_(\d+)_of_21_")
CUDA_RE = re.compile(r"(cuda|unspecified launch failure|cudaerrorlaunchfailure|cublas|device-side assert)", re.IGNORECASE)
ERROR_RE = re.compile(r"(error|exception|runtimeerror|failed|timeout|oom|out of memory)", re.IGNORECASE)


def _read_jsonl(path: Path) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    if not path.exists():
        return rows
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        text = line.strip()
        if not text:
            continue
        try:
            row = json.loads(text)
        except Exception:
            continue
        if isinstance(row, dict):
            rows.append(row)
    return rows


def _discover_episodes(logs_root: Path) -> List[Path]:
    root = logs_root / "btaudit2"
    out: List[Path] = []
    if not root.exists():
        return out
    req = [
        "episode_summary.json",
        "bt.json",
        "status_window_manifest.jsonl",
        "human_rater_evaluation.xlsx",
        "missionengine.jsonl",
    ]
    for ep in sorted(root.glob("*episode_*_of_21_*")):
        if not ep.is_dir():
            continue
        m = EPISODE_RE.search(ep.name)
        if not m:
            continue
        idx = int(m.group(1))
        if idx < 2 or idx > 21:
            continue
        if all((ep / name).exists() for name in req):
            out.append(ep)
    return out


def _audit_missionengine(path: Path) -> Dict[str, Any]:
    response_rows = 0
    response_ok_rows = 0
    response_error_rows = 0
    response_cuda_rows = 0
    empty_response_rows = 0

    for row in _read_jsonl(path):
        if str(row.get("direction", "") or "") != "response":
            continue
        request_kind = str(row.get("request_kind", "") or "")
        if "status_eval" not in request_kind:
            continue

        response_rows += 1
        response_ok = bool(row.get("response_ok", False))
        response_text = str(row.get("response_text", "") or "")
        response_error = str(row.get("response_error", "") or "")
        response_payload = row.get("response") if isinstance(row.get("response"), dict) else {}
        response_error_payload = str(response_payload.get("error", "") or "")

        if response_ok:
            response_ok_rows += 1
        else:
            response_error_rows += 1
            if not response_text.strip():
                empty_response_rows += 1

        haystack = "\n".join([response_text, response_error, response_error_payload])
        if CUDA_RE.search(haystack):
            response_cuda_rows += 1

    ratio = float(response_ok_rows) / float(response_rows) if response_rows > 0 else 0.0
    return {
        "mission_response_rows": response_rows,
        "mission_response_ok_rows": response_ok_rows,
        "mission_response_error_rows": response_error_rows,
        "mission_cuda_error_rows": response_cuda_rows,
        "mission_empty_response_rows": empty_response_rows,
        "mission_response_ok_ratio": ratio,
    }


def _audit_sheet2(path: Path) -> Dict[str, Any]:
    out = {
        "sheet2_exists": False,
        "sheet2_data_rows": 0,
        "sheet2_blank_status_rows": 0,
        "sheet2_error_status_rows": 0,
        "sheet2_unparsed_status_rows": 0,
        "sheet2_error_note_rows": 0,
    }
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Title is more than 31 characters.*", category=UserWarning)
        wb = load_workbook(path, data_only=True)

    if "Sheet2_MissionEngine_Review" not in wb.sheetnames:
        return out

    ws = wb["Sheet2_MissionEngine_Review"]
    out["sheet2_exists"] = True
    if ws.max_row < 9:
        return out

    for r in range(9, ws.max_row + 1):
        source = str(ws.cell(r, 2).value or "").strip()
        if not source:
            continue
        out["sheet2_data_rows"] += 1
        status = str(ws.cell(r, 10).value or "").strip().lower()
        notes = str(ws.cell(r, 13).value or "").strip()
        if not status:
            out["sheet2_blank_status_rows"] += 1
        elif status == "error":
            out["sheet2_error_status_rows"] += 1
        elif status == "unparsed":
            out["sheet2_unparsed_status_rows"] += 1
        if ERROR_RE.search(notes) or CUDA_RE.search(notes):
            out["sheet2_error_note_rows"] += 1

    return out


def _parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Audit btaudit2 mission eval quality and produce rerun manifests")
    ap.add_argument("--logs-root", default=str(Path(__file__).resolve().parents[1] / "data/logs"))
    ap.add_argument("--output-dir", default="", help="Optional output directory. Default: data/logs/btaudit2_eval_audits/<run_tag>")
    ap.add_argument("--min-ok-ratio", type=float, default=1.0)
    return ap.parse_args()


def main() -> None:
    args = _parse_args()
    logs_root = Path(args.logs_root).expanduser().resolve()
    run_tag = time.strftime("%Y%m%dT%H%M%SZ", time.gmtime())
    out_dir = (
        Path(args.output_dir).expanduser().resolve()
        if str(args.output_dir).strip()
        else (logs_root / "btaudit2_eval_audits" / run_tag)
    )
    out_dir.mkdir(parents=True, exist_ok=True)

    episodes = _discover_episodes(logs_root)
    rows: List[Dict[str, Any]] = []
    issue_rows: List[Dict[str, Any]] = []
    cuda_rows: List[Dict[str, Any]] = []

    for ep in episodes:
        mission = _audit_missionengine(ep / "missionengine.jsonl")
        sheet2 = _audit_sheet2(ep / "human_rater_evaluation.xlsx")

        reasons: List[str] = []
        if mission["mission_response_rows"] == 0:
            reasons.append("no_mission_response_rows")
        if mission["mission_response_error_rows"] > 0:
            reasons.append(f"mission_response_error_rows={mission['mission_response_error_rows']}")
        if mission["mission_cuda_error_rows"] > 0:
            reasons.append(f"mission_cuda_error_rows={mission['mission_cuda_error_rows']}")
        if mission["mission_response_ok_ratio"] < float(args.min_ok_ratio):
            reasons.append(f"mission_low_ok_ratio={mission['mission_response_ok_ratio']:.3f}")

        if not sheet2["sheet2_exists"]:
            reasons.append("sheet2_missing")
        elif sheet2["sheet2_data_rows"] == 0:
            reasons.append("sheet2_no_data_rows")
        else:
            if sheet2["sheet2_blank_status_rows"] > 0:
                reasons.append(f"sheet2_blank_status_rows={sheet2['sheet2_blank_status_rows']}")
            if sheet2["sheet2_error_status_rows"] > 0:
                reasons.append(f"sheet2_error_status_rows={sheet2['sheet2_error_status_rows']}")
            if sheet2["sheet2_unparsed_status_rows"] > 0:
                reasons.append(f"sheet2_unparsed_status_rows={sheet2['sheet2_unparsed_status_rows']}")

        issue = len(reasons) > 0
        row = {
            "episode_dir": str(ep),
            **mission,
            **sheet2,
            "issue": issue,
            "reason": "clean" if not issue else ",".join(reasons),
            "has_cuda": bool(mission["mission_cuda_error_rows"] > 0),
        }
        rows.append(row)
        if issue:
            issue_rows.append(row)
        if row["has_cuda"]:
            cuda_rows.append(row)

    audit_json = out_dir / "audit_rows.json"
    issue_manifest = out_dir / "issue_episodes_manifest.txt"
    cuda_manifest = out_dir / "cuda_issue_episodes_manifest.txt"
    summary_json = out_dir / "summary.json"

    audit_json.write_text(json.dumps(rows, indent=2), encoding="utf-8")
    issue_manifest.write_text(
        "\n".join(row["episode_dir"] for row in issue_rows) + ("\n" if issue_rows else ""),
        encoding="utf-8",
    )
    cuda_manifest.write_text(
        "\n".join(row["episode_dir"] for row in cuda_rows) + ("\n" if cuda_rows else ""),
        encoding="utf-8",
    )

    summary = {
        "logs_root": str(logs_root),
        "episodes_checked": len(episodes),
        "episodes_with_issues": len(issue_rows),
        "episodes_with_cuda_issues": len(cuda_rows),
        "min_ok_ratio": float(args.min_ok_ratio),
        "audit_rows_json": str(audit_json),
        "issue_manifest": str(issue_manifest),
        "cuda_issue_manifest": str(cuda_manifest),
    }
    summary_json.write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print("SUMMARY_JSON", summary_json)
    print("EPISODES_CHECKED", len(episodes))
    print("EPISODES_WITH_ISSUES", len(issue_rows))
    print("EPISODES_WITH_CUDA_ISSUES", len(cuda_rows))
    print("ISSUE_MANIFEST", issue_manifest)
    print("CUDA_ISSUE_MANIFEST", cuda_manifest)


if __name__ == "__main__":
    main()
