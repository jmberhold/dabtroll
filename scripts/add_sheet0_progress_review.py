"""Add sheet0_Progress_Review to each workbook in a human-eval schedule.

The script reads a schedule workbook (default: data/logs/human_eval_schedule_final.xlsx),
resolves each linked human_rater_evaluation workbook, and creates/replaces a sheet named
sheet0_Progress_Review using Sheet1 reviewer completion timing plus video duration.
"""

from __future__ import annotations

import argparse
import csv
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Set, Tuple

import cv2
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


HEADER_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass
class NodeRow:
    order: int
    node_id: str
    node_description: str
    success_criteria: str
    completion_sec: Optional[int]
    source_row: int


@dataclass
class WorkbookReport:
    survey_path: str
    video_path: str
    video_duration_seconds: int
    status: str
    sheet0_total_rows: int
    sheet0_data_rows: int
    sheet2_total_rows: int
    sheet2_data_rows: int
    sheet3_total_rows: int
    sheet3_data_rows: int
    sheet4_total_rows: int
    sheet4_data_rows: int
    max_total_rows_sheet0_4: int
    max_data_rows_sheet0_4: int


def _norm_col(name: object) -> str:
    return "".join(ch if str(ch).isalnum() else "_" for ch in str(name or "").strip().lower()).strip("_")


def _find_col_index(headers: Sequence[object], aliases: Sequence[str]) -> Optional[int]:
    norm_to_idx: Dict[str, int] = {}
    for idx, value in enumerate(headers, start=1):
        norm_to_idx[_norm_col(value)] = idx

    for alias in aliases:
        key = _norm_col(alias)
        if key in norm_to_idx:
            return norm_to_idx[key]

    for alias in aliases:
        key = _norm_col(alias)
        for header_key, idx in norm_to_idx.items():
            if key in header_key:
                return idx
    return None


def _parse_mmss_to_seconds(value: object) -> Optional[int]:
    text = str(value or "").strip().strip("'").strip()
    if not text or ":" not in text:
        return None
    parts = text.split(":")
    if len(parts) != 2:
        return None
    try:
        minutes = int(parts[0].strip())
        seconds = int(parts[1].strip())
    except Exception:
        return None
    if minutes < 0 or seconds < 0:
        return None
    return minutes * 60 + seconds


def _to_mmss(seconds: int) -> str:
    sec = max(0, int(seconds))
    return f"{sec // 60}:{sec % 60:02d}"


def _resolve_link(base_dir: Path, raw: object) -> Optional[Path]:
    text = str(raw or "").strip()
    if not text:
        return None
    if text.startswith("file://"):
        return Path(text.replace("file://", "", 1)).expanduser().resolve()
    p = Path(text)
    if p.is_absolute():
        return p
    return (base_dir / p).resolve()


def _find_sheet1(ws_names: Sequence[str], workbook) -> Optional[object]:
    for name in ws_names:
        ws = workbook[name]
        for header_row in (8, 9, 1):
            if header_row > ws.max_row:
                continue
            headers = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
            has_node = _find_col_index(headers, ["node_id"]) is not None
            has_desc = _find_col_index(headers, ["node_description"]) is not None
            has_time = _find_col_index(
                headers,
                ["reviewer_completion_time_m:ss", "reviewer_time_m:ss", "reviewer_completion_time"],
            ) is not None
            if has_node and has_desc and has_time:
                return ws
    return None


def _find_sheet_by_tokens(workbook, token_groups: Sequence[Sequence[str]]) -> Optional[object]:
    for name in workbook.sheetnames:
        ws = workbook[name]
        for header_row in (8, 9, 4, 1):
            if header_row > ws.max_row:
                continue
            headers = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
            found_all = True
            for group in token_groups:
                if _find_col_index(headers, group) is None:
                    found_all = False
                    break
            if found_all:
                return ws
    return None


def _find_header_row(ws, aliases: Sequence[str]) -> Optional[int]:
    for header_row in (8, 9, 4, 1):
        if header_row > ws.max_row:
            continue
        headers = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
        if _find_col_index(headers, aliases) is not None:
            return header_row
    return None


def _sheet_counts(ws, header_aliases: Sequence[str]) -> Tuple[int, int]:
    if ws is None:
        return (0, 0)
    total_rows = int(ws.max_row or 0)
    header_row = _find_header_row(ws, header_aliases)
    if header_row is None:
        return (total_rows, 0)
    return (total_rows, max(0, total_rows - header_row))


def _extract_nodes(sheet1) -> Tuple[int, List[NodeRow]]:
    header_row = None
    node_col = desc_col = crit_col = time_col = order_col = None

    for r in (8, 9, 1):
        if r > sheet1.max_row:
            continue
        headers = [sheet1.cell(row=r, column=c).value for c in range(1, sheet1.max_column + 1)]
        node_col = _find_col_index(headers, ["node_id"])
        desc_col = _find_col_index(headers, ["node_description"])
        crit_col = _find_col_index(headers, ["success_criteria"])
        time_col = _find_col_index(
            headers,
            ["reviewer_completion_time_m:ss", "reviewer_time_m:ss", "reviewer_completion_time"],
        )
        order_col = _find_col_index(headers, ["node_order", "order", "milestone_idx"])
        if node_col is not None and desc_col is not None and time_col is not None:
            header_row = r
            break

    if header_row is None or node_col is None or desc_col is None or time_col is None:
        return (0, [])

    rows: List[NodeRow] = []
    inferred_order = 1
    for r in range(header_row + 1, sheet1.max_row + 1):
        node_id = str(sheet1.cell(row=r, column=node_col).value or "").strip()
        if not node_id:
            continue

        if order_col is not None:
            try:
                order = int(sheet1.cell(row=r, column=order_col).value)
            except Exception:
                order = inferred_order
        else:
            order = inferred_order
        inferred_order += 1

        node_description = str(sheet1.cell(row=r, column=desc_col).value or "").strip()
        success_criteria = str(sheet1.cell(row=r, column=crit_col).value or "").strip() if crit_col else ""
        completion_sec = _parse_mmss_to_seconds(sheet1.cell(row=r, column=time_col).value)

        rows.append(
            NodeRow(
                order=order,
                node_id=node_id,
                node_description=node_description,
                success_criteria=success_criteria,
                completion_sec=completion_sec,
                source_row=r,
            )
        )

    rows.sort(key=lambda x: x.order)
    return (header_row, rows)


def _get_video_duration_seconds(video_path: Optional[Path]) -> int:
    if video_path is None or not video_path.exists():
        return 0
    cap = cv2.VideoCapture(str(video_path))
    if not cap.isOpened():
        cap.release()
        return 0
    fps = float(cap.get(cv2.CAP_PROP_FPS) or 0.0)
    frame_count = float(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0.0)
    cap.release()
    if fps <= 0.0 or frame_count <= 0.0:
        return 0
    return int(math.ceil(frame_count / fps))


def _build_review_timestamps(duration_sec: int, nodes: Sequence[NodeRow]) -> Tuple[List[int], Set[int], bool]:
    timestamps: Set[int] = set()
    completion_event_times: Set[int] = set()

    if duration_sec > 0:
        base_max = duration_sec - (duration_sec % 2)
        if base_max >= 2:
            for sec in range(2, base_max + 1, 2):
                timestamps.add(sec)

        add_final = duration_sec % 2 != 0
        if add_final:
            timestamps.add(duration_sec)
    else:
        add_final = False

    completion_pairs = [(idx, n.completion_sec) for idx, n in enumerate(nodes) if n.completion_sec is not None]

    for i in range(len(completion_pairs) - 1):
        _, prev_sec = completion_pairs[i]
        _, next_sec = completion_pairs[i + 1]
        if next_sec - prev_sec < 4:
            timestamps.add(next_sec)
            completion_event_times.add(next_sec)

    if nodes and nodes[0].completion_sec is not None and nodes[0].completion_sec < 2:
        timestamps.add(nodes[0].completion_sec)
        completion_event_times.add(nodes[0].completion_sec)

    if not timestamps:
        timestamps.add(2 if duration_sec >= 2 else max(0, duration_sec))

    return (sorted(timestamps), completion_event_times, add_final)


def _node_for_time(nodes: Sequence[NodeRow], t_sec: int) -> Optional[NodeRow]:
    if not nodes:
        return None
    idx = 0
    for i, node in enumerate(nodes[:-1]):
        if node.completion_sec is not None and t_sec > node.completion_sec:
            idx = i + 1
    return nodes[min(idx, len(nodes) - 1)]


def _format_sheet0(workbook, nodes: Sequence[NodeRow], duration_sec: int) -> None:
    if "sheet0_Progress_Review" in workbook.sheetnames:
        old = workbook["sheet0_Progress_Review"]
        workbook.remove(old)

    ws = workbook.create_sheet("sheet0_Progress_Review", 0)

    ws["A1"] = "Human Review Sheet 0: Progress Timeline Review"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:L1")

    ws["A2"] = (
        "Rows are sampled every 2s using video time from 0:00, with additional rows for "
        "rapid consecutive node completions (<4s apart)."
    )
    ws.merge_cells("A2:L2")
    ws["A2"].alignment = Alignment(wrap_text=True)

    headers = [
        "review_time_m:ss",
        "review_time_seconds",
        "row_type",
        "node_id",
        "node_description",
        "success_criteria",
        "reviewer_progress",
        "reviewer_progress_half",
        "reviewer_progress_notes",
        "node_completion_observed_YN",
        "sheet1_row_ref",
        "video_duration_seconds",
    ]
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=text)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    timestamps, completion_event_times, has_final_row = _build_review_timestamps(duration_sec, nodes)

    for out_row, sec in enumerate(timestamps, start=5):
        node = _node_for_time(nodes, sec)
        row_type = "interval_2s"
        if sec in completion_event_times:
            row_type = "completion_event"
        if has_final_row and sec == duration_sec:
            row_type = "final_position"

        ws.cell(row=out_row, column=1, value=_to_mmss(sec))
        ws.cell(row=out_row, column=2, value=sec)
        ws.cell(row=out_row, column=3, value=row_type)
        ws.cell(row=out_row, column=4, value=node.node_id if node else "")
        ws.cell(row=out_row, column=5, value=node.node_description if node else "")
        ws.cell(row=out_row, column=6, value=node.success_criteria if node else "")
        ws.cell(row=out_row, column=7, value="")
        ws.cell(row=out_row, column=8, value="")
        ws.cell(row=out_row, column=9, value="")
        ws.cell(row=out_row, column=10, value="")
        ws.cell(row=out_row, column=11, value=node.source_row if node else "")
        ws.cell(row=out_row, column=12, value=duration_sec)

    for r in range(4, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if c not in (1, 2, 3, 10, 11, 12):
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "A": 16,
        "B": 20,
        "C": 20,
        "D": 34,
        "E": 52,
        "F": 52,
        "G": 18,
        "H": 22,
        "I": 48,
        "J": 28,
        "K": 16,
        "L": 22,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:L{max(4, ws.max_row)}"


def _iter_schedule_targets(schedule_path: Path) -> List[Tuple[Path, Optional[Path]]]:
    wb = load_workbook(schedule_path, data_only=False)
    ws = wb.active

    header_row = None
    for r in (6, 1, 7):
        headers = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if _find_col_index(headers, ["survey_link"]) is not None:
            header_row = r
            break
    if header_row is None:
        raise ValueError(f"Could not find survey_link header in schedule: {schedule_path}")

    headers = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
    survey_col = _find_col_index(headers, ["survey_link"])
    video_col = _find_col_index(headers, ["video_link"])
    if survey_col is None:
        raise ValueError(f"Could not find survey_link column in schedule: {schedule_path}")

    base_dir = schedule_path.parent
    targets: List[Tuple[Path, Optional[Path]]] = []
    seen: Set[Path] = set()

    for r in range(header_row + 1, ws.max_row + 1):
        survey_cell = ws.cell(row=r, column=survey_col)
        survey_raw = survey_cell.hyperlink.target if survey_cell.hyperlink else survey_cell.value
        survey_path = _resolve_link(base_dir, survey_raw)
        if survey_path is None:
            continue

        video_path = None
        if video_col is not None:
            video_cell = ws.cell(row=r, column=video_col)
            video_raw = video_cell.hyperlink.target if video_cell.hyperlink else video_cell.value
            video_path = _resolve_link(base_dir, video_raw)

        if survey_path not in seen:
            targets.append((survey_path, video_path))
            seen.add(survey_path)

    return targets


def _write_report_csv(report_path: Path, rows: Sequence[WorkbookReport]) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "survey_path",
        "video_path",
        "video_duration_seconds",
        "status",
        "sheet0_total_rows",
        "sheet0_data_rows",
        "sheet2_total_rows",
        "sheet2_data_rows",
        "sheet3_total_rows",
        "sheet3_data_rows",
        "sheet4_total_rows",
        "sheet4_data_rows",
        "max_total_rows_sheet0_4",
        "max_data_rows_sheet0_4",
    ]
    with report_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "survey_path": row.survey_path,
                    "video_path": row.video_path,
                    "video_duration_seconds": row.video_duration_seconds,
                    "status": row.status,
                    "sheet0_total_rows": row.sheet0_total_rows,
                    "sheet0_data_rows": row.sheet0_data_rows,
                    "sheet2_total_rows": row.sheet2_total_rows,
                    "sheet2_data_rows": row.sheet2_data_rows,
                    "sheet3_total_rows": row.sheet3_total_rows,
                    "sheet3_data_rows": row.sheet3_data_rows,
                    "sheet4_total_rows": row.sheet4_total_rows,
                    "sheet4_data_rows": row.sheet4_data_rows,
                    "max_total_rows_sheet0_4": row.max_total_rows_sheet0_4,
                    "max_data_rows_sheet0_4": row.max_data_rows_sheet0_4,
                }
            )


def add_progress_review_sheets(
    schedule_path: Path,
    dry_run: bool = False,
    report_csv: Optional[Path] = None,
) -> Dict[str, object]:
    targets = _iter_schedule_targets(schedule_path)
    created = 0
    missing_workbooks = 0
    missing_sheet1 = 0
    report_rows: List[WorkbookReport] = []

    global_max_data_rows = 0
    global_max_total_rows = 0
    global_max_data_source = ""
    global_max_total_source = ""

    for survey_path, video_path in targets:
        if not survey_path.exists():
            missing_workbooks += 1
            report_rows.append(
                WorkbookReport(
                    survey_path=str(survey_path),
                    video_path=str(video_path or ""),
                    video_duration_seconds=0,
                    status="missing_workbook",
                    sheet0_total_rows=0,
                    sheet0_data_rows=0,
                    sheet2_total_rows=0,
                    sheet2_data_rows=0,
                    sheet3_total_rows=0,
                    sheet3_data_rows=0,
                    sheet4_total_rows=0,
                    sheet4_data_rows=0,
                    max_total_rows_sheet0_4=0,
                    max_data_rows_sheet0_4=0,
                )
            )
            continue

        wb = load_workbook(survey_path)
        sheet1 = _find_sheet1(wb.sheetnames, wb)
        if sheet1 is None:
            missing_sheet1 += 1
            report_rows.append(
                WorkbookReport(
                    survey_path=str(survey_path),
                    video_path=str(video_path or ""),
                    video_duration_seconds=0,
                    status="missing_sheet1",
                    sheet0_total_rows=0,
                    sheet0_data_rows=0,
                    sheet2_total_rows=0,
                    sheet2_data_rows=0,
                    sheet3_total_rows=0,
                    sheet3_data_rows=0,
                    sheet4_total_rows=0,
                    sheet4_data_rows=0,
                    max_total_rows_sheet0_4=0,
                    max_data_rows_sheet0_4=0,
                )
            )
            continue

        _, nodes = _extract_nodes(sheet1)
        if not nodes:
            missing_sheet1 += 1
            report_rows.append(
                WorkbookReport(
                    survey_path=str(survey_path),
                    video_path=str(video_path or ""),
                    video_duration_seconds=0,
                    status="missing_sheet1_nodes",
                    sheet0_total_rows=0,
                    sheet0_data_rows=0,
                    sheet2_total_rows=0,
                    sheet2_data_rows=0,
                    sheet3_total_rows=0,
                    sheet3_data_rows=0,
                    sheet4_total_rows=0,
                    sheet4_data_rows=0,
                    max_total_rows_sheet0_4=0,
                    max_data_rows_sheet0_4=0,
                )
            )
            continue

        duration_sec = _get_video_duration_seconds(video_path)
        _format_sheet0(wb, nodes, duration_sec)
        if not dry_run:
            wb.save(survey_path)
        created += 1

        sheet0 = wb["sheet0_Progress_Review"] if "sheet0_Progress_Review" in wb.sheetnames else None
        sheet2 = _find_sheet_by_tokens(wb, [["milestone_idx"], ["node_id"], ["model_status"]])
        sheet3 = _find_sheet_by_tokens(wb, [["milestone_idx"], ["node_id"], ["model_status"]])
        # Prefer an explicit sheet4 if present; fallback to token match.
        sheet4 = None
        for name in wb.sheetnames:
            if _norm_col(name).startswith("sheet4"):
                sheet4 = wb[name]
                break
        if sheet4 is None:
            sheet4 = _find_sheet_by_tokens(wb, [["milestone_idx"], ["node_id"], ["model_status"]])

        # Avoid returning the same sheet object for 2/3/4 when token matches are ambiguous.
        explicit2 = None
        explicit3 = None
        for name in wb.sheetnames:
            n = _norm_col(name)
            if n.startswith("sheet2") or n == _norm_col("Sheet2_MissionEngine_Review"):
                explicit2 = wb[name]
            elif n.startswith("sheet3"):
                explicit3 = wb[name]
        if explicit2 is not None:
            sheet2 = explicit2
        if explicit3 is not None:
            sheet3 = explicit3

        s0_total, s0_data = _sheet_counts(sheet0, ["review_time_m:ss", "review_time_seconds"])
        s2_total, s2_data = _sheet_counts(sheet2, ["milestone_idx", "node_id"])
        s3_total, s3_data = _sheet_counts(sheet3, ["milestone_idx", "node_id"])
        s4_total, s4_data = _sheet_counts(sheet4, ["milestone_idx", "node_id"])

        max_total = max(s0_total, s2_total, s3_total, s4_total)
        max_data = max(s0_data, s2_data, s3_data, s4_data)

        if max_total > global_max_total_rows:
            global_max_total_rows = max_total
            global_max_total_source = str(survey_path)
        if max_data > global_max_data_rows:
            global_max_data_rows = max_data
            global_max_data_source = str(survey_path)

        report_rows.append(
            WorkbookReport(
                survey_path=str(survey_path),
                video_path=str(video_path or ""),
                video_duration_seconds=duration_sec,
                status="updated" if not dry_run else "dry_run",
                sheet0_total_rows=s0_total,
                sheet0_data_rows=s0_data,
                sheet2_total_rows=s2_total,
                sheet2_data_rows=s2_data,
                sheet3_total_rows=s3_total,
                sheet3_data_rows=s3_data,
                sheet4_total_rows=s4_total,
                sheet4_data_rows=s4_data,
                max_total_rows_sheet0_4=max_total,
                max_data_rows_sheet0_4=max_data,
            )
        )

    if report_csv is not None:
        _write_report_csv(report_csv, report_rows)

    return {
        "targets": len(targets),
        "updated": created,
        "missing_workbooks": missing_workbooks,
        "missing_sheet1": missing_sheet1,
        "dry_run": bool(dry_run),
        "report_csv": str(report_csv) if report_csv else "",
        "global_max_total_rows_sheet0_4": global_max_total_rows,
        "global_max_total_rows_source": global_max_total_source,
        "global_max_data_rows_sheet0_4": global_max_data_rows,
        "global_max_data_rows_source": global_max_data_source,
    }


def _parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Add sheet0_Progress_Review to human eval workbooks from a schedule file")
    ap.add_argument(
        "--schedule",
        default=str(Path(__file__).resolve().parents[1] / "data/logs/human_eval_schedule_final.xlsx"),
        help="Path to the schedule workbook with survey_link/video_link columns",
    )
    ap.add_argument(
        "--report-csv",
        default="",
        help="Optional CSV output path with per-workbook row counts for sheet0/sheet2/sheet3/sheet4",
    )
    ap.add_argument(
        "--dry-run",
        action="store_true",
        help="Build sheet0 and report counts without saving workbook changes",
    )
    return ap.parse_args()


def main() -> None:
    args = _parse_args()
    schedule_path = Path(args.schedule).expanduser().resolve()
    if not schedule_path.exists():
        raise SystemExit(f"Schedule workbook not found: {schedule_path}")

    report_csv = Path(args.report_csv).expanduser().resolve() if args.report_csv else None
    stats = add_progress_review_sheets(
        schedule_path=schedule_path,
        dry_run=bool(args.dry_run),
        report_csv=report_csv,
    )
    print(stats)


if __name__ == "__main__":
    main()
