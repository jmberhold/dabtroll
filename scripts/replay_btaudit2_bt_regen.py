#!/usr/bin/env python3
"""Regenerate BTs and refresh human node sheets; use only on backed-up datasets."""

from __future__ import annotations

import argparse
import json
import re
import tempfile
import time
import warnings
from copy import copy
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

from dabtroll_bt_pipeline import (
    MissionEngineClient,
    _build_bt_validation,
    _validate_btaudit_bt_structure,
)
from dabtroll_bt_planner import parse_json_from_text
from generate_human_eval_workbook import generate_workbook
from mission_engine import bt_to_graphviz


EPISODE_RE = re.compile(r"_episode_(\d+)_of_21_")
CUDA_RE = re.compile(
    r"(cuda|cudnn|unspecified launch failure|cudaerrorlaunchfailure|cublas|device-side assert|out of memory)",
    re.IGNORECASE,
)
SHEET1_NAME = "Sheet1_BT_Node_Review"


def _read_json(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _append_jsonl(path: Path, payload: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(payload, ensure_ascii=True) + "\n")


def _read_jsonl(path: Path) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    if not path.exists():
        return rows
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        text = line.strip()
        if not text:
            continue
        try:
            row = json.loads(text)
        except Exception:
            continue
        if isinstance(row, dict):
            rows.append(row)
    return rows


def _is_cuda_text(*parts: Any) -> bool:
    text = "\n".join(str(p or "") for p in parts)
    return bool(CUDA_RE.search(text))


def _discover_episodes(logs_root: Path) -> List[Path]:
    root = logs_root / "btaudit2"
    episodes: List[Path] = []
    if not root.exists():
        return episodes
    for ep in sorted(root.glob("*episode_*_of_21_*")):
        if not ep.is_dir():
            continue
        m = EPISODE_RE.search(ep.name)
        if not m:
            continue
        idx = int(m.group(1))
        if idx < 2 or idx > 21:
            continue
        if not (ep / "episode_summary.json").exists():
            continue
        if not (ep / "human_rater_evaluation.xlsx").exists():
            continue
        episodes.append(ep)
    return episodes


def _choose_start_frame(episode_dir: Path) -> Optional[Path]:
    direct = episode_dir / "start_frame.jpg"
    if direct.exists():
        return direct
    for p in sorted((episode_dir / "frames").glob("**/frame_*.jpg")):
        return p
    return None


def _get_task_text(summary: Dict[str, Any]) -> str:
    for key in ("default_task", "high_level_task", "task_text"):
        value = str(summary.get(key, "") or "").strip()
        if value:
            return value
    return ""


def _copy_sheet(src_ws: Any, dst_ws: Any) -> None:
    for row in src_ws.iter_rows():
        for cell in row:
            out = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                out.font = copy(cell.font)
                out.fill = copy(cell.fill)
                out.border = copy(cell.border)
                out.alignment = copy(cell.alignment)
                out.number_format = copy(cell.number_format)
                out.protection = copy(cell.protection)
            if cell.comment is not None:
                out.comment = copy(cell.comment)
            if cell.hyperlink is not None:
                out._hyperlink = copy(cell.hyperlink)

    for key, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[key].width = dim.width
        dst_ws.column_dimensions[key].hidden = dim.hidden

    for idx, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[idx].height = dim.height
        dst_ws.row_dimensions[idx].hidden = dim.hidden

    for merged in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merged))

    dst_ws.freeze_panes = src_ws.freeze_panes
    if src_ws.auto_filter is not None:
        dst_ws.auto_filter.ref = src_ws.auto_filter.ref


def _refresh_sheet1_only(episode_dir: Path) -> None:
    workbook_path = episode_dir / "human_rater_evaluation.xlsx"
    if not workbook_path.exists():
        raise FileNotFoundError(f"workbook missing: {workbook_path}")

    with tempfile.NamedTemporaryFile(prefix="bt_sheet1_", suffix=".xlsx", delete=False) as tmp:
        temp_path = Path(tmp.name)

    try:
        generate_workbook(episode_dir=episode_dir, output_path=temp_path)

        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Title is more than 31 characters.*", category=UserWarning)
            src_wb = load_workbook(temp_path)
            dst_wb = load_workbook(workbook_path)

        if SHEET1_NAME not in src_wb.sheetnames:
            raise RuntimeError(f"generated workbook missing {SHEET1_NAME}")

        src_ws = src_wb[SHEET1_NAME]
        dst_index = 0
        if SHEET1_NAME in dst_wb.sheetnames:
            dst_index = dst_wb.sheetnames.index(SHEET1_NAME)
            old_ws = dst_wb[SHEET1_NAME]
            dst_wb.remove(old_ws)

        dst_ws = dst_wb.create_sheet(title=SHEET1_NAME, index=dst_index)
        _copy_sheet(src_ws, dst_ws)
        dst_wb.save(workbook_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def _update_episode_docs(
    episode_dir: Path,
    bt_json: Dict[str, Any],
    raw_bt_text: str,
    bt_validation: Dict[str, Any],
    bt_variant: str,
    bt_metrics: Dict[str, int],
) -> None:
    bt_json_path = episode_dir / "bt.json"
    bt_raw_path = episode_dir / "qwen_bt_raw.txt"
    bt_validation_path = episode_dir / "bt_validation.json"

    _write_json(bt_json_path, bt_json)
    bt_raw_path.write_text(str(raw_bt_text or ""), encoding="utf-8")
    _write_json(bt_validation_path, bt_validation)
    bt_to_graphviz(bt_json, episode_dir, out_name="bt", fmt="svg")

    summary_path = episode_dir / "episode_summary.json"
    if summary_path.exists():
        summary = _read_json(summary_path)
    else:
        summary = {}

    summary.update(
        {
            "bt_json_path": str(bt_json_path),
            "bt_raw_path": str(bt_raw_path),
            "bt_svg_path": str(episode_dir / "bt.svg"),
            "bt_validation_path": str(bt_validation_path),
            "mission_engine_bt_ok": True,
            "mission_engine_bt_error": "",
            "mission_engine_bt_variant_requested": "btaudit",
            "mission_engine_bt_variant_returned": str(bt_variant or ""),
            "bt_abort_reason": "",
            "bt_action_nodes": int(bt_metrics.get("action", 0)),
            "bt_condition_nodes": int(bt_metrics.get("condition", 0)),
            "bt_control_nodes": int(bt_metrics.get("control", 0)),
            "bt_hash": str(bt_validation.get("bt_hash", "")),
            "bt_depth": int(bt_validation.get("depth", 0)),
            "btaudit_reuse_bt": False,
            "btaudit_source_episode": None,
        }
    )
    _write_json(summary_path, summary)

    manifest_path = episode_dir / "episode_manifest.json"
    if manifest_path.exists():
        manifest = _read_json(manifest_path)
    else:
        manifest = {}
    artifacts = manifest.get("artifacts")
    if not isinstance(artifacts, dict):
        artifacts = {}
    artifacts.update(
        {
            "bt_json": str(bt_json_path),
            "bt_raw": str(bt_raw_path),
            "bt_svg": str(episode_dir / "bt.svg"),
            "bt_validation": str(bt_validation_path),
        }
    )
    manifest["artifacts"] = artifacts
    if isinstance(summary, dict):
        manifest["summary"] = summary
    _write_json(manifest_path, manifest)


def _regenerate_episode_bt(
    *,
    client: MissionEngineClient,
    episode_dir: Path,
    bt_max_new_tokens: int,
) -> Dict[str, Any]:
    started = time.time()
    summary = _read_json(episode_dir / "episode_summary.json")
    mission_name = str(summary.get("mission_name") or episode_dir.name)
    task_text = _get_task_text(summary)
    if not task_text:
        return {
            "ok": False,
            "error": "missing_task_text",
            "is_cuda": False,
            "episode_dir": str(episode_dir),
        }

    start_frame = _choose_start_frame(episode_dir)
    if start_frame is None:
        return {
            "ok": False,
            "error": "missing_start_frame",
            "is_cuda": False,
            "episode_dir": str(episode_dir),
        }

    payload = {
        "mode": "bt",
        "image_path": str(start_frame),
        "frames_b64": [],
        "mission_name": mission_name,
        "task_text": task_text,
        "bt_variant": "btaudit",
        "bt_max_new_tokens": int(bt_max_new_tokens),
    }

    response = client.request(payload, retries=0)
    raw_bt_text = str(response.get("text", "") or "")
    bt_json = parse_json_from_text(raw_bt_text) if response.get("ok") else None
    bt_retry_count = 0
    bt_retry_tokens = 0

    if response.get("ok") and bt_json is None:
        bt_retry_count = 1
        bt_retry_tokens = max(2, int(bt_max_new_tokens) * 2)
        retry_payload = dict(payload)
        retry_payload["bt_max_new_tokens"] = int(bt_retry_tokens)
        retry_response = client.request(retry_payload, retries=0)
        retry_text = str(retry_response.get("text", "") or "")
        retry_bt_json = parse_json_from_text(retry_text) if retry_response.get("ok") else None
        if retry_response.get("ok"):
            response = retry_response
            raw_bt_text = retry_text
            bt_json = retry_bt_json

    error_text = str(response.get("error", "") or "")
    bt_variant = str(response.get("bt_variant", "") or "")
    is_cuda = _is_cuda_text(error_text, raw_bt_text)

    if not response.get("ok"):
        return {
            "ok": False,
            "error": f"mission_engine_request_failed:{error_text}",
            "is_cuda": is_cuda,
            "episode_dir": str(episode_dir),
            "bt_retry_count": int(bt_retry_count),
            "bt_retry_tokens": int(bt_retry_tokens),
        }
    if bt_json is None:
        return {
            "ok": False,
            "error": "mission_engine_bt_invalid_json",
            "is_cuda": is_cuda,
            "episode_dir": str(episode_dir),
            "bt_retry_count": int(bt_retry_count),
            "bt_retry_tokens": int(bt_retry_tokens),
        }

    if str(bt_variant).strip().lower() != "btaudit":
        return {
            "ok": False,
            "error": f"mission_engine_bt_variant_mismatch:{bt_variant}",
            "is_cuda": is_cuda,
            "episode_dir": str(episode_dir),
        }

    valid_bt, validate_reason, metrics = _validate_btaudit_bt_structure(bt_json)
    if not valid_bt:
        return {
            "ok": False,
            "error": str(validate_reason),
            "is_cuda": is_cuda,
            "episode_dir": str(episode_dir),
            "bt_action_nodes": int(metrics.get("action", 0)),
            "bt_condition_nodes": int(metrics.get("condition", 0)),
            "bt_control_nodes": int(metrics.get("control", 0)),
        }

    bt_validation = _build_bt_validation(bt_json, bt_version="dabtroll_bt_v1", btaudit_enabled=True)
    _update_episode_docs(
        episode_dir=episode_dir,
        bt_json=bt_json,
        raw_bt_text=raw_bt_text,
        bt_validation=bt_validation,
        bt_variant=bt_variant,
        bt_metrics=metrics,
    )
    _refresh_sheet1_only(episode_dir)

    return {
        "ok": True,
        "error": "",
        "is_cuda": is_cuda,
        "episode_dir": str(episode_dir),
        "elapsed_s": round(time.time() - started, 3),
        "bt_hash": str(bt_validation.get("bt_hash", "")),
        "bt_depth": int(bt_validation.get("depth", 0)),
        "bt_action_nodes": int(metrics.get("action", 0)),
        "bt_condition_nodes": int(metrics.get("condition", 0)),
        "bt_control_nodes": int(metrics.get("control", 0)),
        "bt_retry_count": int(bt_retry_count),
        "bt_retry_tokens": int(bt_retry_tokens),
    }


def _load_successes(state_jsonl: Path) -> set[str]:
    out: set[str] = set()
    for row in _read_jsonl(state_jsonl):
        if bool(row.get("ok")) and row.get("episode_dir"):
            out.add(str(row.get("episode_dir")))
    return out


def _write_pending_manifest(path: Path, episodes: Sequence[Path], successes: set[str]) -> List[str]:
    pending = [str(ep) for ep in episodes if str(ep) not in successes]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(pending) + ("\n" if pending else ""), encoding="utf-8")
    return pending


def _append_line(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(text + "\n")


def _parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Regenerate personalized BTs for btaudit2 episodes 2-21")
    ap.add_argument("--logs-root", default=str(Path(__file__).resolve().parents[1] / "data/logs"))
    ap.add_argument("--mission-host", default="127.0.0.1")
    ap.add_argument("--mission-port", type=int, default=5560)
    ap.add_argument("--mission-timeout-ms", type=int, default=120000)
    ap.add_argument("--bt-max-new-tokens", type=int, default=1024)
    ap.add_argument("--max-consec-fails", type=int, default=5)
    ap.add_argument("--max-total-fails", type=int, default=25)
    ap.add_argument("--max-consec-cuda-fails", type=int, default=3)
    ap.add_argument("--resume-skip-successes", type=int, default=1)
    ap.add_argument("--run-tag", default="")
    ap.add_argument("--state-jsonl", default="")
    ap.add_argument("--source-manifest", default="", help="Optional manifest of episode directories to process")
    return ap.parse_args()


def _discover_from_manifest(path: Path) -> List[Path]:
    rows: List[Path] = []
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        text = line.strip()
        if not text:
            continue
        p = Path(text)
        if p.exists() and p.is_dir():
            rows.append(p)
    return rows


def _unique_paths(rows: Iterable[Path]) -> List[Path]:
    out: List[Path] = []
    seen: set[str] = set()
    for p in rows:
        key = str(p)
        if key in seen:
            continue
        seen.add(key)
        out.append(p)
    return out


def main() -> None:
    args = _parse_args()
    logs_root = Path(args.logs_root).expanduser().resolve()
    run_tag = str(args.run_tag or time.strftime("%Y%m%dT%H%M%SZ", time.gmtime()))
    run_root = logs_root / "btaudit2_bt_regen_runs" / run_tag
    run_root.mkdir(parents=True, exist_ok=True)

    txt_report = run_root / "run_report.txt"
    error_report = run_root / "error_report.txt"
    json_report = run_root / "run_report.json"
    pending_manifest = run_root / "pending_manifest.txt"

    state_jsonl = Path(args.state_jsonl).expanduser().resolve() if args.state_jsonl else (logs_root / "btaudit2" / "bt_regen_resume_state.jsonl")
    state_jsonl.parent.mkdir(parents=True, exist_ok=True)
    state_jsonl.touch(exist_ok=True)

    if args.source_manifest:
        source_manifest = Path(args.source_manifest).expanduser().resolve()
        if not source_manifest.exists():
            raise SystemExit(f"SOURCE_MANIFEST not found: {source_manifest}")
        discovered = _discover_from_manifest(source_manifest)
    else:
        discovered = _discover_episodes(logs_root)

    discovered = _unique_paths(discovered)
    discovered.sort()

    prior_success = _load_successes(state_jsonl)
    if int(args.resume_skip_successes) == 1:
        targets = [ep for ep in discovered if str(ep) not in prior_success]
    else:
        targets = discovered

    _append_line(txt_report, f"RUN_TAG={run_tag}")
    _append_line(txt_report, f"RUN_ROOT={run_root}")
    _append_line(txt_report, f"STATE_JSONL={state_jsonl}")
    _append_line(txt_report, f"TOTAL_DISCOVERED={len(discovered)}")
    _append_line(txt_report, f"TOTAL_TO_RUN={len(targets)}")
    _append_line(txt_report, f"RESUME_SKIP_SUCCESSES={int(args.resume_skip_successes)}")
    _append_line(txt_report, f"MISSION_HOST={args.mission_host}")
    _append_line(txt_report, f"MISSION_PORT={int(args.mission_port)}")
    _append_line(txt_report, f"MISSION_TIMEOUT_MS={int(args.mission_timeout_ms)}")
    _append_line(txt_report, f"BT_MAX_NEW_TOKENS={int(args.bt_max_new_tokens)}")

    if not targets:
        pending = _write_pending_manifest(pending_manifest, discovered, prior_success)
        payload = {
            "ok": True,
            "stopped_early": False,
            "stop_reason": "",
            "run_rows": [],
            "remaining_count": len(pending),
            "pending_manifest": str(pending_manifest),
            "state_jsonl": str(state_jsonl),
            "txt_report": str(txt_report),
            "error_report": str(error_report),
        }
        _write_json(json_report, payload)
        print(f"JSON_REPORT {json_report}")
        print(f"PENDING_MANIFEST {pending_manifest}")
        print(f"STATE_JSONL {state_jsonl}")
        return

    client = MissionEngineClient(
        host=str(args.mission_host),
        port=int(args.mission_port),
        timeout_ms=int(args.mission_timeout_ms),
    )

    run_rows: List[Dict[str, Any]] = []
    ok_count = 0
    fail_count = 0
    consec_fails = 0
    consec_cuda_fails = 0
    stopped_early = False
    stop_reason = ""

    total = len(targets)
    for idx, ep in enumerate(targets, start=1):
        _append_line(txt_report, f"[{idx}/{total}] RUN {ep}")

        result = _regenerate_episode_bt(
            client=client,
            episode_dir=ep,
            bt_max_new_tokens=int(args.bt_max_new_tokens),
        )
        row = {
            "ts": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
            **result,
        }
        run_rows.append(row)
        _append_jsonl(state_jsonl, row)

        if bool(result.get("ok")):
            ok_count += 1
            consec_fails = 0
            consec_cuda_fails = 0
            _append_line(
                txt_report,
                (
                    "  OK "
                    f"hash={result.get('bt_hash','')} "
                    f"nodes(a/c/ctrl)={result.get('bt_action_nodes',0)}/"
                    f"{result.get('bt_condition_nodes',0)}/{result.get('bt_control_nodes',0)} "
                    f"elapsed_s={result.get('elapsed_s',0)}"
                ),
            )
            continue

        fail_count += 1
        consec_fails += 1
        is_cuda = bool(result.get("is_cuda", False))
        if is_cuda:
            consec_cuda_fails += 1
        else:
            consec_cuda_fails = 0

        _append_line(
            txt_report,
            f"  FAIL error={result.get('error','unknown')} is_cuda={int(is_cuda)}",
        )
        _append_line(error_report, f"[{idx}/{total}] episode={ep}")
        _append_line(error_report, f"error={result.get('error','unknown')}")
        _append_line(error_report, f"is_cuda={int(is_cuda)}")
        _append_line(error_report, "")

        if consec_fails >= int(args.max_consec_fails):
            stopped_early = True
            stop_reason = "max_consecutive_failures_reached"
        elif fail_count >= int(args.max_total_fails):
            stopped_early = True
            stop_reason = "max_total_failures_reached"
        elif consec_cuda_fails >= int(args.max_consec_cuda_fails):
            stopped_early = True
            stop_reason = "max_consecutive_cuda_failures_reached"

        if stopped_early:
            _append_line(txt_report, f"STOP_EARLY reason={stop_reason} at_index={idx}")
            break

    all_success = _load_successes(state_jsonl)
    pending = _write_pending_manifest(pending_manifest, discovered, all_success)

    payload = {
        "ok": True,
        "stopped_early": bool(stopped_early),
        "stop_reason": stop_reason,
        "run_rows_count": len(run_rows),
        "run_ok_count": int(ok_count),
        "run_fail_count": int(fail_count),
        "remaining_count": len(pending),
        "pending_manifest": str(pending_manifest),
        "state_jsonl": str(state_jsonl),
        "txt_report": str(txt_report),
        "error_report": str(error_report),
        "episodes": run_rows,
    }
    _write_json(json_report, payload)

    print(f"JSON_REPORT {json_report}")
    print(f"RUN_OK_COUNT {ok_count}")
    print(f"RUN_FAIL_COUNT {fail_count}")
    print(f"REMAINING_COUNT {len(pending)}")
    print(f"PENDING_MANIFEST {pending_manifest}")
    print(f"STATE_JSONL {state_jsonl}")


if __name__ == "__main__":
    main()
