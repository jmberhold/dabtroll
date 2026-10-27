"""Prepare a portable local dataset candidate; never contact an upload service."""

from __future__ import annotations

import argparse
from collections import Counter
from contextlib import ExitStack
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tarfile
from urllib.parse import urlsplit
import warnings
import zipfile

import pandas as pd
import pyarrow as pa
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

if __package__:
    from scripts.build_btaudit2_visualization_dataset import CURRENT_SCHEMA_COLUMNS, SCHEMA_VERSION
    from scripts.replay_video_utils import _find_primary_video_path, _safe_video_duration_seconds
else:
    from build_btaudit2_visualization_dataset import CURRENT_SCHEMA_COLUMNS, SCHEMA_VERSION
    from replay_video_utils import _find_primary_video_path, _safe_video_duration_seconds


ROOT = Path(__file__).resolve().parents[1]
EPISODES_CSV = "btaudit2_all_episodes_visualization_rows.csv"
OBSERVATIONS_CSV = "btaudit2_all_episodes_visualization_observations.csv"
SENSITIVE_KEY = re.compile(
    r"^(?:api[_-]?key|access[_-]?token|auth(?:orization)?|password|secret|cookie|"
    r"host(?:name)?|.*_host|.*_endpoint|endpoint|base_url|.*_base_url|"
    r"email|reviewer_name|username|creator|last_modified_by)$", re.I,
)
LOCAL_PATH = re.compile(
    r"(?<![\w:/])(?:file://)?/(?:home|Users|mnt|media|tmp|opt|workspace|data|var|root|srv)"
    r"(?:/[^\s\"'<>;,)\]}\\]+)*|(?<![\w:/])(?:file:///)?[A-Za-z]:[\\/][^\s\"'<>;,)\]}]+"
)
URL = re.compile(r"(?:https?|tcp|wss?)://[^\s\"'<>]+")
PRIVATE_ADDRESS = re.compile(
    r"\b(?:127|10)(?:\.\d{1,3}){3}\b|\b192\.168(?:\.\d{1,3}){2}\b|"
    r"\b172\.(?:1[6-9]|2\d|3[01])(?:\.\d{1,3}){2}\b"
)
EMAIL = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")
CREDENTIAL = re.compile(r"\b(?:sk-[A-Za-z0-9_-]{16,}|hf_[A-Za-z0-9]{20,}|"
                        r"AKIA[A-Z0-9]{16}|Bearer\s+[A-Za-z0-9_.-]{16,})")
COHORTS = ("prompt_development", "evaluation")
TABLES = (
    "episodes", "nodes", "human_annotations", "workbook_rows",
    "progress_annotations", "model_observations", "replay_attempts", "frames",
)
DATASET_TAGS = (
    "robotics", "embodied-ai", "robot-simulation", "robot-manipulation", "humanoid-robotics",
    "gr00t", "nvidia-gr00t", "isaac-gr00t", "gr00t-n1.6", "gr00t-n1.6-3b",
    "robocasa", "robocasa-gr1-tabletop-tasks", "gr1", "fourier-gr1", "gr1-tabletop-tasks",
    "robosuite", "mujoco", "tabletop-manipulation", "household-manipulation",
    "kitchen-manipulation", "pick-and-place", "grasping", "drawer-closing", "cabinet-closing",
    "can-to-drawer", "wine-to-cabinet", "cuttingboard-to-pot", "tray-to-plate",
    "vision-language-action", "vla", "vision-language-models", "vlm",
    "behavior-trees", "behavior-tree-auditing", "robot-auditing", "policy-evaluation",
    "vlm-evaluation", "human-evaluation", "human-annotations", "task-completion", "task-progress",
    "success-detection", "failure-detection", "temporal-reasoning", "sequential-tasks",
    "common-evidence-replay", "prompt-evaluation", "qwen", "qwen3-vl", "qwen3.5",
    "qwen3.8", "molmo2", "gemma-4", "luna6", "multimodal",
    "image", "video", "text", "tabular", "timeseries", "csv", "parquet", "datasets", "pandas",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=True, allow_nan=False, default=str)


class Portable:
    def __init__(self, root: Path):
        self.root = root.resolve()
        self.changes: Counter[str] = Counter()

    def path(self, match: re.Match) -> str:
        value = match.group().removeprefix("file://").replace("\\", "/")
        if "/btaudit2/" in value:
            relative = value.split("/btaudit2/", 1)[1]
            if re.match(r"[^/]+_episode_\d+_of_\d+_[^/]+(?:/|$)", relative):
                self.changes["dataset_path"] += 1
                return "episodes/" + relative
            if "/" not in relative and relative.endswith(".csv"):
                self.changes["summary_path"] += 1
                return "summaries/" + relative
        self.changes["local_path"] += 1
        return "[LOCAL_PATH_REDACTED]"

    def url(self, match: re.Match) -> str:
        parsed = urlsplit(match.group())
        host = parsed.hostname or ""
        if (parsed.username or parsed.password or host in {"localhost", "0.0.0.0", "::1"}
                or "." not in host or host.endswith((".local", ".internal"))
                or host.startswith(("127.", "10.", "192.168."))
                or re.match(r"172\.(?:1[6-9]|2\d|3[01])\.", host)):
            self.changes["endpoint"] += 1
            return "[LOCAL_ENDPOINT_REDACTED]"
        return match.group()

    def clean(self, value: object) -> object:
        if isinstance(value, dict):
            result = {}
            for key, item in value.items():
                if SENSITIVE_KEY.fullmatch(str(key)):
                    self.changes["private_field"] += 1
                    result[str(key)] = "[PRIVATE_FIELD_REDACTED]"
                else:
                    result[str(key)] = self.clean(item)
            return result
        if isinstance(value, (list, tuple)):
            return [self.clean(item) for item in value]
        if not isinstance(value, str):
            return None if pd.isna(value) else value
        if CREDENTIAL.search(value):
            raise ValueError("Potential credential detected in release input; review locally before exporting")
        if value.lstrip().startswith(("{", "[")):
            try:
                parsed = json.loads(value)
            except json.JSONDecodeError:
                pass
            else:
                return json_text(self.clean(parsed))
        value = URL.sub(self.url, value)
        if PRIVATE_ADDRESS.search(value):
            self.changes["address"] += 1
            value = PRIVATE_ADDRESS.sub("[LOCAL_ADDRESS_REDACTED]", value)
        value = LOCAL_PATH.sub(self.path, value)
        if EMAIL.search(value):
            self.changes["contact"] += 1
            value = EMAIL.sub("[CONTACT_REDACTED]", value)
        return value

    def frame(self, frame: pd.DataFrame) -> pd.DataFrame:
        result = frame.copy()
        for column in result:
            if SENSITIVE_KEY.fullmatch(str(column)):
                result[column] = "[PRIVATE_FIELD_REDACTED]"
            elif result[column].dtype == object or pd.api.types.is_string_dtype(result[column]):
                result[column] = result[column].map(self.clean)
        return result


class Sources:
    def __init__(self, root: Path):
        self.root = root.resolve()
        self.files: dict[str, dict] = {}

    def add(self, path: Path) -> str:
        if path.is_symlink() or not path.resolve().is_relative_to(self.root):
            raise ValueError(f"Source is a symlink or outside dataset root: {path.name}")
        name = path.relative_to(self.root).as_posix()
        digest = sha256(path)
        previous = self.files.get(name)
        if previous is not None and previous["sha256"] != digest:
            raise ValueError(f"Source changed while preparing release: {name}")
        self.files[name] = {"path": name, "sha256": digest, "bytes": path.stat().st_size}
        return digest

    def verify(self) -> None:
        for name, record in self.files.items():
            if sha256(self.root / name) != record["sha256"]:
                raise ValueError(f"Source changed while preparing release: {name}")


def stage_paper_bundle(source: Path, output: Path, portable: Portable) -> dict:
    """Stage portable paper assets and a standalone ZIP without changing originals."""
    source, output = source.resolve(), output.resolve()
    if source.is_relative_to(output) or output.is_relative_to(source):
        raise ValueError("Paper source and release output must not contain each other")
    target = output / "paper"
    archive = output / "vibtarve_paper.zip"
    if target.exists() or archive.exists():
        raise FileExistsError("Paper assets already exist in the release output")
    if not (source / "manifest.json").is_file():
        raise FileNotFoundError("Paper bundle has no manifest.json")
    sources, before = Sources(source), portable.changes.copy()
    target.mkdir(parents=True)
    for path in sorted(source.rglob("*")):
        if path.is_symlink():
            raise ValueError(f"Paper source is a symlink: {path.relative_to(source)}")
        if not path.is_file():
            continue
        sources.add(path)
        destination = target / path.relative_to(source)
        destination.parent.mkdir(parents=True, exist_ok=True)
        if path.suffix == ".csv":
            frame = pd.read_csv(path, dtype=str, keep_default_na=False)
            portable.frame(frame).to_csv(destination, index=False)
        elif path.suffix == ".json":
            write_json(destination, portable.clean(json.loads(path.read_text())))
        elif path.suffix in {".html", ".md", ".tex"}:
            destination.write_text(str(portable.clean(path.read_text())), encoding="utf-8")
        elif path.suffix in {".png", ".pdf"}:
            shutil.copyfile(path, destination)
        else:
            raise ValueError(f"Unsupported paper asset: {path.relative_to(source)}")

    for path in sorted(target.rglob("manifest.json")):
        catalog = json.loads(path.read_text())
        original = json.loads((source / path.relative_to(target)).read_text())
        copied = Path(catalog["source"]["copy"])
        if copied.is_absolute() or ".." in copied.parts:
            raise ValueError("Paper catalog source copy must be a relative contained path")
        original_digest = original["source"]["sha256"]
        if sha256(source / path.parent.relative_to(target) / copied) != original_digest:
            raise ValueError(f"Paper catalog has a stale source CSV: {path.relative_to(target)}")
        digest = sha256(path.parent / copied)
        catalog["source"]["original_sha256"] = original_digest
        catalog["source"]["sha256"] = digest
        for entry in catalog["entries"]:
            if "source_sha256" in entry:
                if entry["source_sha256"] != original_digest:
                    raise ValueError(f"Paper catalog contains stale entries: {path.relative_to(target)}")
                entry["original_source_sha256"] = original_digest
                entry["source_sha256"] = digest
            for name in entry["files"].values():
                relative = Path(name)
                if relative.is_absolute() or ".." in relative.parts or not (path.parent / relative).is_file():
                    raise ValueError(f"Invalid or missing paper asset reference: {name}")
        write_json(path, catalog)
        for name in ("index.html", "README.md"):
            page = path.parent / name
            if page.is_file():
                text = page.read_text().replace(original_digest, digest)
                text = text.replace("Full original episode CSV", "Portable episode CSV")
                text = text.replace("Full unchanged source CSV", "Portable source CSV")
                page.write_text(text, encoding="utf-8")
    sources.verify()
    records = [
        {**record, "released_sha256": sha256(target / name)}
        for name, record in sorted(sources.files.items())
    ]
    write_json(target / "release_provenance.json", {
        "transform": "Paths/contacts sanitized; CSV shapes preserved; PNG/PDF bytes unchanged",
        "catalog_hashes": "Released CSV copies; original source hashes retained separately",
        "privacy_transformations": dict(portable.changes - before),
        "source_files_verified_unchanged": len(records), "files": records,
    })
    with (target / "SHA256SUMS").open("w") as handle:
        for path in sorted(target.rglob("*")):
            if path.is_file() and path != target / "SHA256SUMS":
                handle.write(f"{sha256(path)}  {path.relative_to(target).as_posix()}\n")
    shutil.make_archive(str(archive.with_suffix("")), "zip", root_dir=output, base_dir="paper")
    return {"directory": "paper", "archive": archive.name, "source_file_count": len(records),
            "file_count": len(records) + 2, "archive_sha256": sha256(archive)}


def probe_video(path: Path) -> float | None:
    import cv2

    duration = _safe_video_duration_seconds(path)
    if duration is None or duration <= 0:
        return None
    capture = cv2.VideoCapture(str(path))
    try:
        readable, _ = capture.read()
    finally:
        capture.release()
    return duration if readable else None


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=True, allow_nan=False, default=str)
                    + "\n", encoding="utf-8")


def write_table(output: Path, name: str, rows: list[dict], portable: Portable) -> dict:
    frame = portable.frame(pd.DataFrame(rows))
    if frame.empty:
        frame = pd.DataFrame(columns=["episode_dir", "cohort"])
    directory = output / "data" / name
    directory.mkdir(parents=True, exist_ok=True)
    schema = pa.Schema.from_pandas(frame, preserve_index=False)
    for cohort in COHORTS:
        part = frame[frame["cohort"].eq(cohort)]
        part.to_parquet(directory / f"{cohort}.parquet", index=False, schema=schema)
    if name in {"human_annotations", "progress_annotations"}:
        frame.to_csv(directory / "all.csv", index=False)
    return {"rows": len(frame), "cohorts": frame["cohort"].value_counts().to_dict(),
            "columns": list(frame.columns)}


def workbook_rows(path: Path, episode: dict, sources: Sources, portable: Portable) -> tuple[list, list]:
    digest = sources.add(path)
    raw_rows, annotations = [], []
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        for sheet in workbook:
            header = None
            for number, row in enumerate(sheet.iter_rows(), 1):
                cells = {
                    get_column_letter(cell.column): {
                        "value": portable.clean(cell.value),
                        **({"hyperlink": portable.clean(cell.hyperlink.target)} if cell.hyperlink else {}),
                    }
                    for cell in row if cell.value is not None or cell.hyperlink
                }
                if not cells:
                    continue
                raw_rows.append({
                    "episode_dir": episode["episode_dir"], "cohort": episode["cohort"],
                    "sheet": sheet.title, "row_number": number, "cells_json": json_text(cells),
                    "source_workbook_sha256": digest,
                })
                values = [cell.value for cell in row]
                normalized = [str(v).strip().lower() if v is not None else "" for v in values]
                if sheet.title.lower().startswith("sheet1") and "node_id" in normalized:
                    header = normalized
                    continue
                if header is not None:
                    record = {key: portable.clean(value) for key, value in zip(header, values) if key}
                    if record.get("node_id"):
                        annotations.append({
                            "episode_dir": episode["episode_dir"], "cohort": episode["cohort"],
                            "confirmed_review": episode["human_reviewed"],
                            "source_sheet": sheet.title, "source_row": number,
                            "source_workbook_sha256": digest, **record,
                        })
    finally:
        workbook.close()
    return raw_rows, annotations


def stage_episode_access(source: Path, output: Path) -> dict:
    """Expose already-sanitized reviews and verified frame pixels by episode."""
    source, output = source.resolve(), output.resolve()
    if (output / "episode_catalog.csv").exists():
        raise FileExistsError("Episode access exports already exist")

    def table(name):
        return pd.concat(
            pd.read_parquet(source / "data" / name / f"{cohort}.parquet")
            for cohort in COHORTS
        ).reset_index(drop=True)

    episodes, annotations, rows, frames = (
        table(name) for name in ("episodes", "human_annotations", "workbook_rows", "frames")
    )
    names = set(episodes.episode_dir)
    if episodes.episode_dir.duplicated().any() or any(
        Path(name).name != name or name in {".", ".."} for name in names
    ):
        raise ValueError("Invalid episode access keys")
    for frame in (annotations, rows, frames):
        if not set(frame.episode_dir).issubset(names):
            raise ValueError("Episode access table references unknown episodes")
    sheets = rows[rows.sheet.str.lower().str.startswith("sheet1")]
    if set(sheets.episode_dir) != names or set(frames.episode_dir) != names:
        raise ValueError("Every episode must have Sheet1 rows and saved frames")
    if frames.duplicated(["archive_path", "member_path"]).any():
        raise ValueError("Duplicate frame archive members")
    for row in frames.itertuples():
        member = Path(row.member_path)
        expected = Path("episodes") / row.episode_dir / "frames" / row.camera
        archive = Path(row.archive_path)
        if (member.parent != expected or ".." in member.parts or member.is_absolute()
                or archive.is_absolute() or ".." in archive.parts
                or not archive.parts or archive.parts[:2] != ("media", "frames")):
            raise ValueError("Unsafe episode frame reference")
    for name in names:
        target = output / "episodes" / name
        for filename in ("human_review.xlsx", "sheet1_human_review.csv", "frames.zip",
                         "frame_index.csv", "README.md"):
            if (target / filename).exists():
                raise FileExistsError(target / filename)
    counts = Counter()
    for archive_name, group in frames.groupby("archive_path", sort=True):
        expected = group.set_index("member_path").to_dict("index")
        with tarfile.open(source / archive_name, "r") as archive, ExitStack() as stack:
            packages = {}
            for member in archive:
                if member.name not in expected or not member.isfile():
                    raise ValueError(f"Unexpected frame member: {archive_name}/{member.name}")
                record = expected.pop(member.name)
                stream = archive.extractfile(member)
                if stream is None:
                    raise ValueError(f"Cannot read frame: {member.name}")
                payload = stream.read()
                if (len(payload) != record["bytes"]
                        or hashlib.sha256(payload).hexdigest() != record["sha256"]):
                    raise ValueError(f"Frame checksum mismatch: {member.name}")
                name = record["episode_dir"]
                if name not in packages:
                    target = output / "episodes" / name / "frames.zip"
                    target.parent.mkdir(parents=True, exist_ok=True)
                    packages[name] = stack.enter_context(zipfile.ZipFile(target, "a", compression=zipfile.ZIP_STORED))
                relative = Path(member.name).relative_to(Path("episodes") / name).as_posix()
                info = zipfile.ZipInfo(relative)
                info.external_attr = 0o644 << 16
                if relative in packages[name].namelist():
                    raise ValueError(f"Duplicate episode frame: {relative}")
                packages[name].writestr(info, payload)
                counts[name] += 1
            if expected:
                raise ValueError(f"Missing indexed frame members: {archive_name}")
    records = []
    for episode in episodes.sort_values(["cohort", "scenario_id", "episode_dir"]).to_dict("records"):
        name = episode["episode_dir"]
        target = output / "episodes" / name
        sheet = sheets[sheets.episode_dir.eq(name)].sort_values("row_number")
        if (sheet.sheet.nunique() != 1 or sheet.row_number.duplicated().any()
                or sheet.source_workbook_sha256.nunique() != 1):
            raise ValueError(f"Ambiguous Sheet1 rows: {name}")
        workbook = Workbook()
        workbook.properties.creator = "VIBTARVE"
        workbook.properties.lastModifiedBy = "VIBTARVE"
        page = workbook.active
        page.title = str(sheet.iloc[0]["sheet"])
        for row in sheet.itertuples():
            for column, cell in json.loads(row.cells_json).items():
                copied = page[f"{column}{row.row_number}"]
                copied.value = cell["value"]
                if isinstance(cell["value"], str):
                    copied.data_type = "s"
                if cell.get("hyperlink"):
                    copied.hyperlink = cell["hyperlink"]
        workbook.save(target / "human_review.xlsx")
        workbook.close()
        reviewed = annotations[annotations.episode_dir.eq(name)].copy()
        if reviewed.empty or not reviewed.confirmed_review.eq(bool(episode["human_reviewed"])).all():
            raise ValueError(f"Inconsistent confirmed review flags: {name}")
        reviewed.to_csv(target / "sheet1_human_review.csv", index=False)
        index = frames[frames.episode_dir.eq(name)].copy()
        index["zip_member_path"] = index.member_path.map(
            lambda value: Path(value).relative_to(Path("episodes") / name).as_posix()
        )
        index.to_csv(target / "frame_index.csv", index=False)
        status = "Completed human review" if episode["human_reviewed"] else "Unconfirmed / available review template"
        video = episode.get("video_path")
        video = video if isinstance(video, str) and video else ""
        if video and (Path(video).is_absolute() or ".." in Path(video).parts
                      or not (source / video).is_file()
                      or not Path(video).is_relative_to(Path("episodes") / name)):
            raise ValueError(f"Invalid episode video: {name}")
        record = {
            "episode_dir": name, "cohort": episode["cohort"], "scenario_id": episode["scenario_id"],
            "human_reviewed": bool(episode["human_reviewed"]), "review_status": status,
            "review_basis": episode["review_basis"], "sheet1_node_rows": len(reviewed),
            "source_workbook_sha256": str(sheet.iloc[0].source_workbook_sha256),
            "sheet1_xlsx": f"episodes/{name}/human_review.xlsx",
            "sheet1_csv": f"episodes/{name}/sheet1_human_review.csv",
            "frames_zip": f"episodes/{name}/frames.zip",
            "frame_index": f"episodes/{name}/frame_index.csv",
            "frame_count": counts[name], "cameras": ", ".join(sorted(index.camera.unique())),
            "video_path": video, "video_status": episode["video_status"],
        }
        video_link = f"[Video]({Path(video).relative_to(Path('episodes') / name).as_posix()})" if video else "Video unavailable"
        (target / "README.md").write_text(
            f"# {name}\n\n**{status}** ({episode['review_basis']}). "
            "Completed review does not mean successful task completion. Blank completion "
            "times in unconfirmed templates are not negative labels.\n\n"
            "- [Human review: Sheet1 XLSX](human_review.xlsx)\n"
            "- [Human review: Sheet1 CSV](sheet1_human_review.csv)\n"
            f"- [All {counts[name]} saved frames, both cameras](frames.zip)\n"
            "- [Frame index and original pixel SHA256 hashes](frame_index.csv)\n"
            f"- {video_link}\n\n"
            "The XLSX is a reconstructed, sanitized Sheet1-only copy: saved values, "
            "cell coordinates and hyperlinks are retained; formulas are inert text. "
            "Original workbook properties, styling and other sheets are not included. "
            "The CSV includes node rows with review confirmation and source provenance. "
            "Extract frames.zip here to create frames/<camera>/frame_<step>.jpg; "
            "the two recorded views are not interchangeable. Verify member hashes with "
            "frame_index.csv. Source workbook SHA256: "
            f"`{record['source_workbook_sha256']}`.\n\n"
            "To contribute, retain episode/node IDs and the source revision, record your "
            "reviewer revision and evidence in a separate copy, and submit a GitHub issue "
            "or pull request describing the annotation change. Do not overwrite another "
            "reviewer's judgments or treat a downloaded copy as an automatically accepted "
            "annotation. Never publish private reviewer information.\n",
            encoding="utf-8",
        )
        records.append(record)
    catalog = pd.DataFrame(records)
    catalog.to_csv(output / "episode_catalog.csv", index=False)
    lines = [
        "# Episode catalog: human reviews, frames and videos\n\n",
        "Each episode has a Sheet1-only XLSX, node-row CSV, and frame ZIP containing both "
        "camera views. Completed-review status is explicit; it does not mean task success. "
        "Unconfirmed templates must not be scored as completed annotations. "
        "The ZIPs contain unchanged image bytes, not newly sampled video frames.\n\n",
        "[Download the full episode catalog](episode_catalog.csv)\n\n",
    ]
    for cohort in COHORTS:
        part = catalog[catalog.cohort.eq(cohort)]
        for confirmed in (True, False):
            subset = part[part.human_reviewed.eq(confirmed)]
            label = "Completed reviews" if confirmed else "Unconfirmed reviews / templates"
            lines += [f"## {cohort}: {label} ({len(subset)})\n\n",
                      "| Episode | Sheet1 | Frames | Video |\n|---|---|---|---|\n"]
            for row in subset.to_dict("records"):
                video = f"[MP4]({row['video_path']})" if row["video_path"] else "Unavailable"
                lines.append(
                    f"| [{row['episode_dir']}](episodes/{row['episode_dir']}/README.md) "
                    f"| [XLSX]({row['sheet1_xlsx']}) / [CSV]({row['sheet1_csv']}) "
                    f"| [ZIP ({row['frame_count']})]({row['frames_zip']}) | {video} |\n"
                )
            lines.append("\n")
    (output / "EPISODES.md").write_text("".join(lines), encoding="utf-8")
    return {
        "catalog": "episode_catalog.csv", "browse": "EPISODES.md",
        "episode_count": len(records), "frame_count": int(sum(counts.values())),
        "sheet1_node_rows": len(annotations),
        "completed_reviews": catalog[catalog.human_reviewed].groupby("cohort").size().to_dict(),
        "transform": "Reconstructed sanitized Sheet1 XLSX/CSV and per-episode ZIPs from existing released tables and verified frame shards; original pixel bytes unchanged",
    }


def dataset_card(manifest: dict) -> str:
    tags = "\n".join(f"- {tag}" for tag in DATASET_TAGS)
    configs = "\n".join(
        f"- config_name: {name}\n" + ("  default: true\n" if name == "episodes" else "")
        + "  data_files:\n"
        + "".join(f"  - split: {cohort}\n    path: data/{name}/{cohort}.parquet\n"
                  for cohort in COHORTS if manifest["tables"][name]["cohorts"].get(cohort, 0))
        for name in TABLES if manifest["tables"][name]["rows"]
    )
    return f"""---
license: cc-by-4.0
language:
- en
task_categories:
- robotics
tags:
{tags}
configs:
{configs}
---
# VIBTARVE

VLM-Initiated Behavior Tree Auditing of Robot VLA Execution.

RoboCasa GR1 tabletop task recordings of NVIDIA Isaac GR00T N1.6 execution,
human-review annotations, and VLM auditor outputs for robot simulation,
behavior-tree auditing, and vision-language-action policy evaluation.
The simulation stack uses RoboCasa GR1 tabletop tasks, robosuite, and MuJoCo.

**Local release candidate, not an uploaded or published release.** Target:
https://huggingface.co/datasets/{manifest['repository']}.
The maintainer confirmed redistribution rights and reviewer consent for this candidate.
That confirmation is recorded, not an independent legal review of upstream assets.
CC BY 4.0 applies to the released dataset; it does not relicense upstream software,
model weights, simulator assets, or the implementation repository.

## Scope

{manifest['episode_count']} recordings: {manifest['cohorts'].get('prompt_development', 0)}
prompt-development recordings and {manifest['cohorts'].get('evaluation', 0)} evaluation
recordings. These are related scenario/run trials, **not a random train/test split**.
{manifest['reviewed_evaluation_count']} evaluation recordings have confirmed completed
human review. Workbook presence and empty completion cells alone do not establish review.
Human completion flags are null outside confirmed available node review.
Evaluation confirmation comes from the completed review schedule. Prompt-development
confirmation comes from the explicitly reviewed episode-1 analysis CSV; the evaluation
schedule does not cover that separate cohort. `review_basis` records the distinction.

## Layout and loading

Start with the [episode catalog](EPISODES.md) to find completed human reviews,
Sheet1-only XLSX/CSV downloads, per-episode frame ZIPs containing both cameras,
and videos. [episode_catalog.csv](episode_catalog.csv) provides the same links
and explicit review-confirmation status for programmatic use.

```python
from datasets import load_dataset
episodes = load_dataset("{manifest['repository']}", "episodes", split="evaluation")
annotations = load_dataset("{manifest['repository']}", "human_annotations", split="evaluation")
```

These examples become remote-loadable only after a separately approved upload.
Locally use `pandas.read_parquet("data/episodes/evaluation.parquet")`.
Each configuration is a separate relational table, joined by `episode_dir`
and, where applicable, `node_id`, sheet/row, model name, or frame step.
Only nonempty cohorts are advertised as loadable splits; in this snapshot progress
annotations have an evaluation split but no prompt-development split.

- `data/`: Parquet tables, plus CSV copies of normalized human/progress annotations.
- `summaries/`: **all source-root BTAudit2 CSVs**, with original filenames and
  column layouts. Paths/contact fields are sanitized. `provenance/summary_exports.json`
  distinguishes current canonical exports from historical analyses.
- `episodes/`: portable BT/recording metadata, simulator-state evidence, and
  readable videos. Absolute source paths are rebased or explicitly redacted.
- `media/frames/`: uncompressed TAR shards of both saved camera views. JPEG pixels
  are unchanged. `frames` identifies each archive/member, step, camera, and checksum.
  Archive members are relative `episodes/<episode>/frames/<camera>/frame_<step>.jpg`.
  Extract chosen shards into the release root to use rebased image paths.
  Extract only verified archives and reject absolute/traversal member paths.
- `provenance/prompts/`: saved replay prompt snapshots, content-addressed after
  sanitization. Original and released hashes are distinguished.
- `examples/`: representative scene galleries, not independent annotations.
- `provenance/`: source hashes, cohort lists, attempt configurations, exclusions,
  export policy, and privacy transformations. Raw credentials and debug logs are excluded.

The `workbook_rows` table preserves **all nonempty cells and hyperlinks in every
worksheet**, including historical model sheets, Sheet0 progress, Sheet1 human review,
headers, notes, and formulas as text. Coordinates and workbook hashes retain provenance;
workbook styling, macros, document properties, and original XLSX files are not distributed.
`human_annotations` additionally exposes Sheet1 node rows as columns. Partial/unconfirmed
annotations remain available but are not automatically scored as ground truth.

## Measurement and limitations

Primary selected auditor observations come from the canonical schema-{SCHEMA_VERSION}
tables; historical conditions remain in workbook rows and replay-attempt records.
Synthetic BT milestones are labeled separately from mission-engine responses.
Synthetic milestones do not supply scoring judgments. Comparisons use ordinal
Node 1, Node 2, etc., preserving raw IDs and full model-specific contracts in
`model_node_contracts_json`. Observation rows retain `node_position`,
`node_position_basis`, and `model_node_count`. Only shared positions are scored;
extra positions are not comparable, not false incomplete. Unreached positions
inside known contracts remain incomplete; missing model episodes are not negatives.
Whole-mission accuracy excludes unequal full node counts. Positional agreement
does not establish semantic equivalence of historical BT versions.
Explicitly labeled legacy prefix inferences are not verified historical ID bindings.
Unresolved contracts remain not comparable.

Machine/simulation time differs from review-video time. Luna's derived video times use
frame-index/duration scaling; original timestamps and timing-basis metadata are retained.
Historical Qwen3.8 uses eight-frame contact sheets; Luna uses ten separate images.
These are not controlled input-matched model comparisons. The historical Qwen3-VL
clock discrepancy remains unresolved; do not silently recalibrate it.

Unreadable videos are excluded from playable media, not from episode/node/annotation
tables or available saved frames. Consult `provenance/quality_issues.json` for explicit
exclusions and workbook-lock warnings. Video checks measure duration and decode the first
frame; they are not a guarantee that every encoded frame is undamaged.

Human progress is distinct from binary node completion. Original model explanations
are claims, not independent visual ground truth. Missing labels remain missing.
Review all flagged artifacts before public distribution. Automated privacy checks do
not prove that free-text notes contain no personal information.

## Attribution and reproducibility

Authors: J. Mark Berhold, Bruce A. Cox, Neil C. Ranly, and Nicholas J. Yielding,
Air Force Institute of Technology. No publication venue, DOI, or paper acceptance is claimed.
Implementation and upstream attribution:
https://github.com/jmberhold/dabtroll and its `docs/research_attribution.md`.
The source code/worktree state is recorded in `release_manifest.json`; a dirty worktree
is not an immutable software release. No model weights or upstream simulator assets
are redistributed as standalone assets.

Use `SHA256SUMS` to verify released files. Per-frame hashes verify archive contents.
Source fingerprints describe the immutable on-disk inputs, not unsaved reviewer edits.
This exporter never uploads or changes repository visibility.
"""


def validate_release(output: Path) -> dict:
    manifest = json.loads((output / "release_manifest.json").read_text())
    episodes = pd.concat(pd.read_parquet(output / "data/episodes" / f"{c}.parquet") for c in COHORTS)
    if len(episodes) != manifest["episode_count"] or episodes.episode_dir.duplicated().any():
        raise ValueError("Release episode count or uniqueness mismatch")
    frames = pd.concat(pd.read_parquet(output / "data/frames" / f"{c}.parquet") for c in COHORTS)
    checked = 0
    for archive, records in frames.groupby("archive_path", sort=False):
        expected = records.set_index("member_path").sha256.to_dict()
        with tarfile.open(output / archive, "r") as handle:
            for member in handle:
                path = Path(member.name)
                if path.is_absolute() or ".." in path.parts or not member.isfile():
                    raise ValueError(f"Unsafe frame archive member: {archive}")
                stream = handle.extractfile(member)
                if stream is None or member.name not in expected:
                    raise ValueError(f"Unexpected frame member: {archive}")
                if hashlib.sha256(stream.read()).hexdigest() != expected.pop(member.name):
                    raise ValueError(f"Frame checksum mismatch: {archive}")
                checked += 1
        if expected:
            raise ValueError(f"Missing archived frames: {archive}")
    if checked != manifest["frame_count"]:
        raise ValueError("Frame index count mismatch")
    for table, info in manifest["tables"].items():
        parts = [pd.read_parquet(output / "data" / table / f"{c}.parquet") for c in COHORTS]
        if sum(map(len, parts)) != info["rows"]:
            raise ValueError(f"Table count mismatch: {table}")
        for part in parts:
            if not set(part.episode_dir).issubset(set(episodes.episode_dir)):
                raise ValueError(f"Unknown episode reference: {table}")
            for column in part.select_dtypes(include=["object", "string"]):
                for value in part[column].dropna():
                    if isinstance(value, str) and (LOCAL_PATH.search(value) or CREDENTIAL.search(value)):
                        raise ValueError(f"Private path or potential credential remains: {table}/{column}")
    for video in episodes.video_path.dropna():
        if video and not (output / video).is_file():
            raise ValueError("Video reference does not exist")
    return {"status": "passed", "episode_count": len(episodes),
            "verified_frame_members": checked, "table_counts_verified": True,
            "automated_text_checks_passed": True}


def prepare_release(root: Path, output: Path, *, expected_count: int = 488,
                    rights_confirmed: bool = False, paper_bundle: Path | None = None,
                    shard_bytes: int = 128 * 1024 * 1024) -> dict:
    if not rights_confirmed:
        raise ValueError("Explicit confirmation of redistribution rights and reviewer consent is required")
    root, output = root.resolve(), output.resolve()
    if output == root or output.is_relative_to(root):
        raise ValueError("Release output must be outside the immutable source dataset")
    if output.exists():
        raise FileExistsError("Release output already exists; choose a new candidate directory")
    sources, portable = Sources(root), Portable(root)
    sources.add(root / EPISODES_CSV)
    episodes = pd.read_csv(root / EPISODES_CSV, low_memory=False)
    if (not CURRENT_SCHEMA_COLUMNS.issubset(episodes.columns)
            or not episodes.visualization_schema_version.eq(SCHEMA_VERSION).all()):
        raise ValueError("Refresh canonical visualization data before preparing a release")
    if len(episodes) != expected_count or episodes.episode_dir.duplicated().any():
        raise ValueError("Unexpected source cohort count or duplicate episodes")
    known = set(episodes.episode_dir)
    if any(Path(name).name != name or name in {".", ".."} for name in known):
        raise ValueError("Unsafe source episode directory")
    tuning_reviewed = set()
    tuning_csv = root / "btaudit2_molmo2_episode1_rows.csv"
    if tuning_csv.is_file():
        sources.add(tuning_csv)
        tuning = pd.read_csv(tuning_csv, low_memory=False)
        if tuning.episode_dir.duplicated().any():
            raise ValueError("Duplicate prompt-development review records")
        tuning_reviewed = set(tuning.loc[tuning.human_reviewed.eq(True), "episode_dir"])
        if not tuning_reviewed.issubset(set(episodes.loc[episodes.episode_index.eq(1), "episode_dir"])):
            raise ValueError("Prompt-development review confirmation references unknown recordings")
    output.mkdir(parents=True)
    marker = output / "INCOMPLETE"
    marker.write_text("Do not upload: release preparation or validation has not finished.\n")
    tables: dict[str, list] = {name: [] for name in TABLES}
    issues, summary_exports = [], []
    archive_states: dict[tuple[str, str], tuple] = {}
    with ExitStack() as stack:
        for index, source in enumerate(episodes.to_dict("records"), 1):
            name = source["episode_dir"]
            episode = root / name
            cohort = "prompt_development" if int(source["episode_index"]) == 1 else "evaluation"
            reviewed = (name in tuning_reviewed if cohort == "prompt_development"
                        else str(source["schedule_complete_original"]).lower() == "true")
            summary_path = episode / "episode_summary.json"
            sources.add(summary_path)
            summary = json.loads(summary_path.read_text())
            record = {
                "episode_dir": name, "cohort": cohort, "scenario_id": source["scenario_id"],
                "run_tag": source["run_tag"], "episode_index": int(source["episode_index"]),
                "seed": summary.get("seed"), "high_level_task": source["high_level_task"],
                "human_reviewed": reviewed, "review_status": "confirmed" if reviewed else "unconfirmed",
                "review_basis": "episode1_analysis_csv" if cohort == "prompt_development"
                else "evaluation_schedule_complete",
                "sim_success_ever": source["sim_success_episode"],
                "sim_success_final": source["sim_success_final"],
                "bt_semantic_rating": source["bt_semantic_rating"],
                "episode_path": f"episodes/{name}", "bt_path": f"episodes/{name}/bt.json",
                "video_path": None, "video_status": "unavailable", "video_duration_s": None,
                "outer_step_seconds": summary.get("outer_step_seconds", 0.4),
                "steps_executed": summary.get("steps_executed"),
                "luna6_time_calibration_json": source["luna6_time_calibration_json"],
                "model_node_contracts_json": source["model_node_contracts_json"],
            }
            for lock in sorted([*episode.glob(".~lock.*"), *episode.glob("~$*")]):
                issues.append({"episode_dir": name, "kind": "workbook_lock_present",
                               "detail": "On-disk snapshot only; unsaved edits are not included."})
                warnings.warn(f"Workbook lock present in {name}; snapshot includes saved bytes only",
                              RuntimeWarning, stacklevel=2)
            for filename in ("bt.json", "bt_validation.json", "episode_summary.json",
                             "episode_manifest.json", "episode_outcome.json", "qwen_bt_raw.txt"):
                path = episode / filename
                if not path.is_file():
                    raise FileNotFoundError(path)
                sources.add(path)
                target = output / "episodes" / name / filename
                target.parent.mkdir(parents=True, exist_ok=True)
                if path.suffix == ".json":
                    write_json(target, portable.clean(json.loads(path.read_text())))
                else:
                    target.write_text(str(portable.clean(path.read_text())), encoding="utf-8")
            for filename in ("states_per_frame.jsonl", "btstatus.jsonl"):
                path = episode / filename
                sources.add(path)
                with path.open() as handle, (output / "episodes" / name / filename).open("w") as target:
                    for line in handle:
                        if line.strip():
                            target.write(json_text(portable.clean(json.loads(line))) + "\n")
            for node in json.loads(source["required_nodes_json"]):
                available = reviewed and node["human_node_available"]
                tables["nodes"].append({
                    "episode_dir": name, "cohort": cohort, **node,
                    "human_complete": node["human_complete"] if available else None,
                    "review_status": ("reviewed_complete" if node["human_complete"] else "reviewed_incomplete")
                    if available else "unconfirmed_or_missing",
                })
            for point in json.loads(source["reviewer_progress_points_json"]):
                tables["progress_annotations"].append({
                    "episode_dir": name, "cohort": cohort, "confirmed_review": reviewed, **point,
                })
            raw, annotations = workbook_rows(episode / "human_rater_evaluation.xlsx",
                                             record, sources, portable)
            tables["workbook_rows"].extend(raw)
            tables["human_annotations"].extend(annotations)
            for config_path in sorted(episode.glob("*/replay_config.json")):
                sources.add(config_path)
                config = json.loads(config_path.read_text())
                snapshot = config_path.parent / "status_prompt_source.txt"
                prompt_path, original_prompt_sha = None, None
                if snapshot.is_file():
                    original_prompt_sha = sources.add(snapshot)
                    content = str(portable.clean(snapshot.read_text())).encode()
                    digest = hashlib.sha256(content).hexdigest()
                    prompt_path = f"provenance/prompts/{digest}.txt"
                    (output / "provenance/prompts").mkdir(parents=True, exist_ok=True)
                    (output / prompt_path).write_bytes(content)
                quality = config_path.parent / "response_quality_summary.json"
                tables["replay_attempts"].append({
                    "episode_dir": name, "cohort": cohort, "attempt": config_path.parent.name,
                    "review_sheet": config.get("review_sheet_name"), "model": config.get("model"),
                    "config_json": json_text(portable.clean(config)),
                    "prompt_path": prompt_path, "source_prompt_sha256": original_prompt_sha,
                    "snapshot_matches_config": original_prompt_sha == config.get("status_prompt_sha256")
                    if original_prompt_sha else False,
                    "quality_record_present": quality.is_file(),
                    "raw_attempt_logs": "retained_local_not_distributed",
                })
            video = _find_primary_video_path(episode)
            if video is not None:
                sources.add(video)
                duration = probe_video(video)
                if duration is not None:
                    relative = f"episodes/{name}/video/{video.name}"
                    target = output / relative
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(video, target)
                    record.update(video_path=relative, video_status="readable",
                                  video_duration_s=duration)
                else:
                    issues.append({"episode_dir": name, "kind": "unreadable_video",
                                   "detail": "Duration or first frame unreadable; video excluded, episode retained."})
            else:
                issues.append({"episode_dir": name, "kind": "missing_video",
                               "detail": "Video unavailable; episode retained."})
            tables["episodes"].append(record)
            image_count = 0
            for frame in sorted((episode / "frames").glob("*/*")):
                if frame.suffix.lower() not in {".jpg", ".jpeg", ".png"}:
                    continue
                digest = sources.add(frame)
                camera = frame.parent.name
                key = (cohort, camera)
                state = archive_states.get(key)
                size = frame.stat().st_size
                if state is None or state[0].offset + size + 1024 > shard_bytes:
                    part = state[2] + 1 if state is not None else 0
                    if state is not None:
                        state[0].close()
                    relative = f"media/frames/{cohort}/{camera}/part-{part:04}.tar"
                    (output / relative).parent.mkdir(parents=True, exist_ok=True)
                    handle = stack.enter_context(tarfile.open(output / relative, "w"))
                    state = (handle, relative, part)
                    archive_states[key] = state
                member_name = f"episodes/{name}/frames/{camera}/{frame.name}"
                member = tarfile.TarInfo(member_name)
                member.size, member.mode, member.mtime = size, 0o644, 0
                with frame.open("rb") as handle:
                    state[0].addfile(member, handle)
                tables["frames"].append({
                    "episode_dir": name, "cohort": cohort, "camera": camera,
                    "frame_step": int(frame.stem.split("_")[-1]),
                    "archive_path": state[1], "member_path": member_name,
                    "sha256": digest, "bytes": size,
                })
                image_count += 1
            if not image_count:
                raise ValueError(f"Episode has no saved frames: {name}")
            if index % 50 == 0 or index == len(episodes):
                print(f"Prepared {index}/{len(episodes)} recordings", flush=True)
    sources.add(root / OBSERVATIONS_CSV)
    observations = pd.read_csv(root / OBSERVATIONS_CSV, low_memory=False)
    if not set(observations.episode_dir).issubset(known):
        raise ValueError("Canonical observations reference unknown episodes")
    observations["cohort"] = observations.episode_index.map(
        lambda value: "prompt_development" if int(value) == 1 else "evaluation")
    review_flags = {row["episode_dir"]: row["human_reviewed"] for row in tables["episodes"]}
    observations["human_reviewed"] = observations.episode_dir.map(review_flags)
    tables["model_observations"] = observations.to_dict("records")
    table_info = {name: write_table(output, name, rows, portable) for name, rows in tables.items()}
    episode_access = stage_episode_access(output, output)
    print("Normalizing summary CSV exports", flush=True)
    (output / "summaries").mkdir()
    for path in sorted(root.glob("*.csv")):
        digest = sources.add(path)
        frame = pd.read_csv(path, dtype=str, keep_default_na=False)
        target = output / "summaries" / path.name
        portable.frame(frame).to_csv(target, index=False)
        summary_exports.append({
            "path": f"summaries/{path.name}", "rows": len(frame), "columns": len(frame.columns),
            "source_sha256": digest, "released_sha256": sha256(target),
            "status": "canonical_current" if path.name.startswith((
                "btaudit2_all_episodes_visualization_", "btaudit2_episodes_2_21_visualization_",
            )) else "historical_analysis_not_primary_ground_truth",
            "transform": "Path/contact sanitization; source filenames, column order and row counts preserved",
        })
    paper_assets = None
    if paper_bundle is not None:
        catalog = json.loads((paper_bundle / "manifest.json").read_text())
        if catalog["source"]["sha256"] != sha256(root / "btaudit2_episodes_2_21_visualization_rows.csv"):
            raise ValueError("Example galleries are stale relative to the canonical evaluation CSV")
        (output / "examples").mkdir()
        for entry in catalog["entries"]:
            if entry["label"].startswith("c") and entry["kind"] == "figure":
                for suffix in ("png", "pdf"):
                    path = paper_bundle / entry["files"][suffix]
                    shutil.copyfile(path, output / "examples" / path.name)
        paper_assets = stage_paper_bundle(paper_bundle, output, portable)
    write_json(output / "provenance/quality_issues.json", issues)
    write_json(output / "provenance/summary_exports.json", summary_exports)
    write_json(output / "provenance/privacy_transformations.json", dict(portable.changes))
    (output / "provenance").mkdir(exist_ok=True)
    with (output / "provenance/source_files.jsonl").open("w") as handle:
        for record in sorted(sources.files.values(), key=lambda value: value["path"]):
            handle.write(json_text(record) + "\n")
    for cohort in COHORTS:
        names = sorted(row["episode_dir"] for row in tables["episodes"] if row["cohort"] == cohort)
        (output / "provenance" / f"{cohort}_episodes.txt").write_text("".join(n + "\n" for n in names))
    manifest = {
        "schema": "vibtarve_release_v1", "status": "local_candidate_not_uploaded",
        "repository": "jmberhold/VIBTARVE", "license": "cc-by-4.0",
        "rights_and_reviewer_consent": "confirmed_by_maintainer",
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "source_schema_version": SCHEMA_VERSION,
        "source_episodes_sha256": sources.files[EPISODES_CSV]["sha256"],
        "episode_count": len(episodes),
        "cohorts": dict(Counter(row["cohort"] for row in tables["episodes"])),
        "reviewed_evaluation_count": sum(r["human_reviewed"] and r["cohort"] == "evaluation"
                                         for r in tables["episodes"]),
        "frame_count": len(tables["frames"]),
        "readable_videos": sum(r["video_status"] == "readable" for r in tables["episodes"]),
        "quality_issue_counts": dict(Counter(issue["kind"] for issue in issues)),
        "summary_csv_count": len(summary_exports), "tables": table_info,
        "paper_bundle": paper_assets,
        "episode_access": episode_access,
        "originals_modified": False, "upload_performed": False,
        "source_code": "https://github.com/jmberhold/dabtroll",
        "code_state": "local_worktree_not_an_immutable_software_release",
        "excluded": ["original XLSX containers/properties", "raw request/debug/console logs",
                     "resume/pending/lock files", "unreadable video bytes", "model weights"],
    }
    manifest["code_commit"] = subprocess.run(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, check=True, capture_output=True, text=True,
    ).stdout.strip()
    manifest["code_worktree_dirty"] = bool(subprocess.run(
        ["git", "status", "--porcelain"], cwd=ROOT, check=True, capture_output=True, text=True,
    ).stdout.strip())
    write_json(output / "provenance/code_files.json", {
        name: sha256(ROOT / name) for name in (
            "scripts/prepare_vibtarve_release.py", "scripts/build_btaudit2_visualization_dataset.py",
            "scripts/replay_video_utils.py", "scripts/notebook_qwen_compare_utils.py",
            "requirements-analysis.txt",
        )
    })
    write_json(output / "release_manifest.json", manifest)
    (output / "README.md").write_text(dataset_card(manifest), encoding="utf-8")
    (output / "LICENSE").write_text(
        "VIBTARVE dataset\n\n"
        "This released dataset is licensed under the Creative Commons Attribution 4.0 "
        "International license (CC BY 4.0).\n"
        "License terms: https://creativecommons.org/licenses/by/4.0/legalcode\n\n"
        "Attribution: J. Mark Berhold, Bruce A. Cox, Neil C. Ranly, Nicholas J. Yielding.\n"
        "Dataset: https://huggingface.co/datasets/jmberhold/VIBTARVE\n\n"
        "Upstream software, model weights, simulator assets, and the implementation "
        "repository retain their own terms; this notice does not relicense them.\n",
        encoding="utf-8")
    (output / ".gitattributes").write_text(
        "*.parquet filter=lfs diff=lfs merge=lfs -text\n"
        "*.tar filter=lfs diff=lfs merge=lfs -text\n"
        "*.mp4 filter=lfs diff=lfs merge=lfs -text\n"
        "*.pdf filter=lfs diff=lfs merge=lfs -text\n"
        "*.zip filter=lfs diff=lfs merge=lfs -text\n"
        "*.xlsx filter=lfs diff=lfs merge=lfs -text\n"
        "*.csv filter=lfs diff=lfs merge=lfs -text\n", encoding="utf-8")
    print("Verifying source immutability and staged media/table integrity", flush=True)
    sources.verify()
    report = validate_release(output)
    report["source_files_verified_unchanged"] = len(sources.files)
    write_json(output / "validation.json", report)
    with (output / "SHA256SUMS").open("w") as handle:
        for path in sorted(output.rglob("*")):
            if path.is_file() and path.name not in {"SHA256SUMS", "INCOMPLETE"}:
                handle.write(f"{sha256(path)}  {path.relative_to(output).as_posix()}\n")
    marker.unlink()
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset-root", type=Path, default=ROOT / "data/logs/btaudit2")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--expected-count", type=int, default=488)
    parser.add_argument("--rights-confirmed", action="store_true")
    parser.add_argument("--paper-bundle", type=Path)
    args = parser.parse_args()
    manifest = prepare_release(args.dataset_root, args.output_dir,
                               expected_count=args.expected_count,
                               rights_confirmed=args.rights_confirmed,
                               paper_bundle=args.paper_bundle)
    print(json.dumps({k: v for k, v in manifest.items() if k != "tables"}, indent=2))


if __name__ == "__main__":
    main()
